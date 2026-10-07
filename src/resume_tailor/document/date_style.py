"""The date style of the user's original resume, so tailored dates read the same.

Experience dates are stored as ``YYYY-MM`` and printed by `render.format_range`. The
output used to be hard-coded ("Jul 2026 - Present"); a resume that writes "July 2026 –
Present" would then change on every tailor. This reads the month style and separator off
the uploaded export once and the renderer reproduces them. No match (or no export)
leaves the legacy output unchanged. No LLM.
"""

from __future__ import annotations

import re
from collections import Counter
from functools import lru_cache
from pathlib import Path
from typing import Literal, TypedDict

MonthStyle = Literal["abbr", "full", "abbr_dot", "numeric"]

_MONTHS = (
    "January", "February", "March", "April", "May", "June", "July", "August",
    "September", "October", "November", "December",
)  # fmt: skip
_NAME = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|June?|July?|Aug(?:ust)?|Sept?(?:ember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\.?"
_POINT = rf"(?:{_NAME}\s+\d{{4}}|\d{{1,2}}/\d{{4}})"
_RANGE = re.compile(
    rf"(?P<start>{_POINT})(?P<sep>\s*(?:-|–|—|to)\s*)(?:{_POINT}|Present|Current|Now)",
    re.IGNORECASE,
)
#: Months an AP-style "Jan." leaves spelled out.
_SPELLED = {3, 4, 5, 6, 7}


class DateStyle(TypedDict):
    month: MonthStyle
    separator: str


#: What the renderer printed before styles were read (`Mon YYYY - Mon YYYY`).
DEFAULT: DateStyle = {"month": "abbr", "separator": " - "}


def _month_style(point: str) -> MonthStyle:
    if "/" in point:
        return "numeric"
    word = point.split()[0]
    if word.endswith("."):
        return "abbr_dot"
    return "full" if word.casefold() in {name.casefold() for name in _MONTHS} and len(word) > 3 else "abbr"


def detect_in_texts(texts: list[str]) -> DateStyle | None:
    """The most common ``(month style, separator)`` among the date ranges in ``texts``."""
    seen: Counter[tuple[MonthStyle, str]] = Counter()
    for text in texts:
        for match in _RANGE.finditer(text):
            seen[(_month_style(match.group("start")), match.group("sep"))] += 1
    if not seen:
        return None
    (month, separator), _count = seen.most_common(1)[0]
    return {"month": month, "separator": separator}


@lru_cache(maxsize=8)
def _detect_cached(path: str, mtime_ns: int) -> DateStyle | None:
    from docx import Document  # noqa: PLC0415

    try:
        doc = Document(path)
    except Exception:  # noqa: BLE001 - an unreadable export just means the legacy style
        return None
    texts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(paragraph.text for paragraph in cell.paragraphs)
    return detect_in_texts(texts)


def detect(path: Path) -> DateStyle | None:
    """The date style of the export at ``path``; None when missing or no ranges found."""
    try:
        return _detect_cached(str(path), path.stat().st_mtime_ns)
    except OSError:
        return None


def format_month(year: int, month: int, style: DateStyle) -> str:
    kind = style["month"]
    if kind == "numeric":
        return f"{month:02d}/{year}"
    full = _MONTHS[month - 1]
    if kind == "full":
        return f"{full} {year}"
    if kind == "abbr_dot":
        if month in _SPELLED:
            return f"{full} {year}"
        short = "Sept" if month == 9 else full[:3]
        return f"{short}. {year}"
    return f"{full[:3]} {year}"
