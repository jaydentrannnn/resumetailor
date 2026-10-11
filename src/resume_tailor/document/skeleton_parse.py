"""Read an edited layout file (`template_skeleton.py`) back into a template profile.

The «tokens» say exactly which field sits where, so nothing about the fields is guessed:
each token is replaced with ordinary sample text (the person's own first entries where
they exist), and the profile's spans point at those replacements. The analyzer still
runs on the result, but only for what tokens cannot say — spacing chrome, table layout —
and with every heading's kind forced from its bookmark. The result is an `AnalyzeResult`
the existing wizard (map → preview → install) takes as-is, plus the realised bytes that
become the new baseline. Comments are stripped; no LLM is involved.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import docx
from docx.oxml.ns import qn

from ..content.data import MasterResume
from . import (
    docx_text,
    profile_validation,
    skeleton_tokens,
    template_analyze,
    template_tagging,
)
from .analysis_types import AnalyzeResult, Issue
from .skeleton_tokens import TOKEN_RE, Token
from .template_profile import (
    CharSpan,
    ContactMapping,
    ContactSlot,
    CustomFieldSlot,
    DetectedSection,
    EducationMapping,
    EnabledSections,
    ExperienceMapping,
    HeaderFieldMapping,
    HeadingPrototype,
    ListMapping,
    OptionalSpan,
    ProjectsMapping,
    SkillsMapping,
    TemplateProfile,
)
from .template_skeleton import BOOKMARK_PREFIX, strip_bookmarks


@dataclass
class _Found:
    pid: int
    start: int
    end: int
    token: Token | None = None
    custom: str | None = None  # a custom line's label
    sample: str = ""
    span: CharSpan | None = None  # where the sample sits after realisation


@dataclass
class ParsedSkeleton:
    raw: bytes
    result: AnalyzeResult
    issues: list[Issue] = field(default_factory=list)


def _strip_comments(doc) -> None:
    """Remove every comment anchor and the comments part: the realised file becomes a
    baseline, and the layout notes must not ride along into it."""
    body = doc.element.body
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(body.iter(qn(tag))):
            el.getparent().remove(el)
    for ref in list(body.iter(qn("w:commentReference"))):
        run = ref.getparent()
        run.remove(ref)
        if not [c for c in run if c.tag != qn("w:rPr")]:
            run.getparent().remove(run)
    for r_id, rel in list(doc.part.rels.items()):
        if "comments" in rel.reltype.rsplit("/", 1)[-1].lower():
            doc.part.drop_rel(r_id)


def _bookmarked(doc, paragraphs) -> dict[str, int]:
    """Bookmark name -> paragraph id, for the skeleton's `rt_sec_*` heading marks. A mark
    Word moved between paragraphs is credited to the next paragraph."""
    ids = {p._p: i for i, p in enumerate(paragraphs)}
    out: dict[str, int] = {}
    for mark in doc.element.body.iter(qn("w:bookmarkStart")):
        name = mark.get(qn("w:name")) or ""
        if not name.startswith(BOOKMARK_PREFIX):
            continue
        holder = mark.getparent()
        if holder.tag != qn("w:p"):
            holder = mark.getnext()
            while holder is not None and holder.tag != qn("w:p"):
                holder = holder.getnext()
        if holder is not None and holder in ids:
            out[name] = ids[holder]
    return out


def _samples(resume: MasterResume | None, styled: dict[str, str]) -> dict[tuple[str, str], str]:
    """Sample text per (kind, field): the person's own first entry where there is one —
    the realised file is the new baseline, so it should read like their resume — and
    dates in the template's own style (`styled`, from the skeleton's meta)."""
    out = {(t.kind, t.field): t.sample for t in skeleton_tokens.TOKENS}
    for key, value in styled.items():
        kind, _, name = key.partition(".")
        if (kind, name) in out and value.strip():
            out[(kind, name)] = value
    if resume is None:
        return out
    c = resume.contact
    for name in ("name", "email", "phone", "location"):
        value = getattr(c, name).strip()
        if value:
            out[("name" if name == "name" else "contact", name)] = value
    pick = {
        ("experience", "company"): lambda e: e.company, ("experience", "title"): lambda e: e.title,
        ("experience", "location"): lambda e: e.location,
        ("experience", "bullet"): lambda e: e.bullets[0].text if e.bullets else "",
        ("education", "school"): lambda e: e.school, ("education", "degree"): lambda e: e.degree,
        ("education", "location"): lambda e: e.location,
        ("project", "name"): lambda e: e.name,
        ("project", "bullet"): lambda e: e.bullets[0].text if e.bullets else "",
    }  # fmt: skip
    firsts = {"experience": resume.experience, "education": resume.education,
              "project": resume.projects}  # fmt: skip
    for (kind, name), get in pick.items():
        entries = firsts[kind]
        value = get(entries[0]).strip() if entries else ""
        if value and "\t" not in value:
            out[(kind, name)] = value
    return out


def _find_tokens(paragraphs, issues: list[Issue]) -> list[_Found]:
    found: list[_Found] = []
    for pid, paragraph in enumerate(paragraphs):
        for m in TOKEN_RE.finditer(paragraph.text):
            name = m.group(1) or m.group(2)
            hit = skeleton_tokens.lookup(name)
            if hit is None:
                issues.append(Issue(
                    code="layout_unknown_token",
                    message=f"«{name}» isn't a field the app knows; it will print as written.",
                ))  # fmt: skip
                continue
            item = _Found(pid, m.start(), m.end())
            if isinstance(hit, str):
                item.custom = hit
            else:
                item.token = hit
            found.append(item)
    return found


def _realise(paragraphs, found: list[_Found], samples, custom_text: dict[str, str]) -> None:
    """Replace every token with its sample, recording where each sample lands."""
    by_pid: dict[int, list[_Found]] = {}
    for item in found:
        by_pid.setdefault(item.pid, []).append(item)
    for pid, items in by_pid.items():
        items.sort(key=lambda f: f.start)
        for item in items:
            if item.custom is not None:
                item.sample = custom_text.get(item.custom.casefold(), item.custom)
            else:
                item.sample = samples[(item.token.kind, item.token.field)]
        shift = 0
        for item in items:
            start = item.start + shift
            item.span = CharSpan(paragraph_id=pid, start=start, end=start + len(item.sample))
            shift += len(item.sample) - (item.end - item.start)
        for item in reversed(items):
            template_tagging.replace_span_with_tag(
                paragraphs[pid], CharSpan(paragraph_id=pid, start=item.start, end=item.end),
                item.sample,
            )


def parse_skeleton(raw: bytes, meta: dict, resume: MasterResume | None = None) -> ParsedSkeleton:
    """Read the edited layout file `raw` (made with `meta`, `Skeleton.meta()`)."""
    doc = docx.Document(io.BytesIO(raw))
    _strip_comments(doc)
    paragraphs = [p for p, _loc in docx_text.iter_document_paragraphs(doc)]
    in_table = [loc is not None for _p, loc in docx_text.iter_document_paragraphs(doc)]
    issues: list[Issue] = []
    marks = _bookmarked(doc, paragraphs)
    strip_bookmarks(doc)
    headings = _locate_headings(meta, marks, paragraphs, issues)
    found = _find_tokens(paragraphs, issues)
    custom_text = {
        c["label"].casefold(): c.get("text") or c["label"] for c in meta.get("custom", [])
    }
    _realise(paragraphs, found, _samples(resume, meta.get("samples", {})), custom_text)

    buffer = io.BytesIO()
    doc.save(buffer)
    realised = buffer.getvalue()
    overrides = {pid: skeleton_tokens.ANALYZER_KIND[s["kind"]] for pid, s in headings}
    result = template_analyze.analyze_docx(raw=realised, overrides=overrides)
    builder = _ProfileBuilder(result, headings, found, meta, paragraphs, in_table, issues)
    profile = builder.build()
    if profile is not None:
        paras = {p.id: p for p in result.paragraphs}
        issues += profile_validation._ProfileValidator(profile, paras).run()
    blocking = any(i.blocking for i in issues) or profile is None
    result = result.model_copy(
        update={
            "suggested_profile": profile,
            "issues": issues,
            "unmapped": [],
            "ready": not blocking,
        }
    )
    return ParsedSkeleton(raw=realised, result=result, issues=issues)


def _locate_headings(meta, marks, paragraphs, issues) -> list[tuple[int, dict]]:
    """(paragraph id, meta section) for each section still in the file, by bookmark,
    else by its heading text."""
    out: list[tuple[int, dict]] = []
    for section in meta.get("sections", []):
        pid = marks.get(section["bookmark"])
        if pid is None:
            title = section["title"].strip().casefold()
            pid = next(
                (i for i, p in enumerate(paragraphs) if p.text.strip().casefold() == title), None
            )
        if pid is None:
            issues.append(Issue(
                code="layout_section_missing",
                message=f"The {section['title']} heading is gone, so that section won't print. "
                "Undo the deletion in Word, or keep it if you meant to drop the section.",
            ))  # fmt: skip
            continue
        out.append((pid, section))
    out.sort(key=lambda t: t[0])
    return out


_REQUIRED = {"experience": "bullet", "education": "degree", "project": "bullet",
             "skills": "label", "list": "item"}  # fmt: skip
_TOKEN_LABEL = {(t.kind, t.field): t.label for t in skeleton_tokens.TOKENS}


class _ProfileBuilder:
    """Turns token positions into a `TemplateProfile`, on top of the analyzer's view of
    the realised file for what tokens do not say (spacing, normalisation, layout)."""

    def __init__(self, result, headings, found, meta, paragraphs, in_table, issues):
        self.result, self.headings, self.found = result, headings, found
        self.meta, self.paragraphs, self.in_table, self.issues = (
            meta, paragraphs, in_table, issues,
        )

    def _block(self, code: str, message: str) -> None:
        self.issues.append(Issue(code=code, message=message, blocking=True))

    def _zones(self) -> tuple[list[_Found], dict[int, list[_Found]]]:
        """Tokens above the first heading, and the tokens under each heading."""
        bounds = [pid for pid, _ in self.headings]
        top: list[_Found] = []
        under: dict[int, list[_Found]] = {pid: [] for pid in bounds}
        for item in self.found:
            owner = max((b for b in bounds if b < item.pid), default=None)
            (top if owner is None else under[owner]).append(item)
        return top, under

    def build(self) -> TemplateProfile | None:
        base = self.result.suggested_profile
        top, under = self._zones()
        update: dict = {"source_sha256": self.result.source_sha256}
        update.update(self._header(top))
        enabled: dict[str, bool] = {k: False for k in EnabledSections.model_fields}
        sections: list[DetectedSection] = []
        seen: set[str] = set()
        for pid, section in self.headings:
            kind = section["kind"]
            sections.append(DetectedSection(
                id=section["id"], title=self.paragraphs[pid].text.strip() or section["title"],
                kind=kind, heading_paragraph_id=pid,
            ))  # fmt: skip
            tokens = under[pid]
            wrong = [t for t in tokens if t.token is None or t.token.kind != kind]
            for t in wrong:
                name = t.token.label if t.token else f"custom: {t.custom}"
                self._block("layout_token_misplaced", f"«{name}» sits under {section['title']}, "
                            "where it doesn't belong. Move it back to its own section.")
            if kind in seen:
                continue
            seen.add(kind)
            mapping = self._mapping(kind, pid, [t for t in tokens if t not in wrong], section)
            if mapping is not None:
                key = {"project": "projects", "list": "list_section"}.get(kind, kind)
                update[key] = mapping
                enabled[key] = True
        if not sections:
            self._block("layout_no_sections", "No section headings are left in the layout file.")
            return None
        update.update(
            enabled=EnabledSections(**enabled),
            section_mode="generic",
            sections=sections,
            heading_prototype=HeadingPrototype(paragraph_id=sections[0].heading_paragraph_id),
            paragraph_count=len(self.paragraphs),
            dropped_paragraph_ids=[],
            warnings=[],
        )
        tabled = any(self.in_table[t.span.paragraph_id] for t in self.found if t.token)
        update["layout"] = "table" if tabled else "paragraph"
        data = (base.model_dump() if base is not None else {}) | {
            k: (v.model_dump() if hasattr(v, "model_dump") else
                [x.model_dump() if hasattr(x, "model_dump") else x for x in v]
                if isinstance(v, list) else v)
            for k, v in update.items()
        }  # fmt: skip
        for kind_key in ("experience", "education", "projects", "skills", "list_section"):
            if kind_key not in update:
                data[kind_key] = None
        try:
            return TemplateProfile.model_validate(data)
        except ValueError as exc:
            self._block("layout_profile_invalid", f"The layout file couldn't be read: {exc}")
            return None

    # -- header ------------------------------------------------------------------------

    def _header(self, top: list[_Found]) -> dict:
        update: dict = {"name_in_header": False, "contact_in_header": False}
        names = [t for t in top if t.token is not None and t.token.kind == "name"]
        if not names:
            self._block("layout_name_missing", "«Name» was deleted, so your name won't print.")
        else:
            update["name_paragraph_id"] = names[0].span.paragraph_id
        contact = [t for t in top if t.token is not None and t.token.kind == "contact"]
        update["contact"] = self._contact(contact)
        stray = [t for t in top if t.token is not None and t.token.kind not in ("name", "contact")]
        for t in stray:
            self._block("layout_token_misplaced", f"«{t.token.label}» is above the first "
                        "section heading. Move it back into its section.")  # fmt: skip
        update["custom_fields"] = self._custom([t for t in top if t.custom is not None])
        for t in (t for t in self.found if t.custom is not None and t not in top):
            self._block("layout_custom_misplaced", f"«custom: {t.custom}» must sit above "
                        "the first section heading.")  # fmt: skip
        return update

    def _contact(self, tokens: list[_Found]) -> ContactMapping | None:
        if not tokens:
            return None
        by_pid: dict[int, list[_Found]] = {}
        for t in tokens:
            by_pid.setdefault(t.span.paragraph_id, []).append(t)
        slots = []
        for pid, items in sorted(by_pid.items()):
            items.sort(key=lambda t: t.span.start)
            text = self.paragraphs[pid].text
            sep = text[items[0].span.end : items[1].span.start] if len(items) > 1 else " • "
            slots.append(ContactSlot(
                paragraph_id=pid, fields=[t.token.field for t in items], separator=sep or " • ",
            ))  # fmt: skip
        fields = [f for s in slots for f in s.fields]
        if len(fields) != len(set(fields)):
            self._block("layout_contact_twice", "A contact field appears twice in the header.")
            return None
        first = slots[0]
        return ContactMapping(
            paragraph_id=first.paragraph_id,
            field_order=fields,
            slots=slots if len(slots) > 1 else [],
            separator=first.separator,
        )

    def _custom(self, tokens: list[_Found]) -> list[CustomFieldSlot]:
        known = {c["label"].casefold(): c for c in self.meta.get("custom", [])}
        slots: list[CustomFieldSlot] = []
        taken: set[str] = set()
        for t in tokens:
            if self.paragraphs[t.pid].text.strip() != t.sample.strip():
                self._block("layout_custom_shared", f"«custom: {t.custom}» must be on a line "
                            "of its own.")  # fmt: skip
                continue
            prior = known.get(t.custom.casefold())
            key = prior["key"] if prior else _slug(t.custom, taken)
            taken.add(key)
            slots.append(CustomFieldSlot(
                key=key, label=t.custom, paragraph_id=t.span.paragraph_id, text=t.sample,
            ))  # fmt: skip
        return slots

    # -- sections ----------------------------------------------------------------------

    def _mapping(self, kind: str, heading: int, tokens: list[_Found], section: dict):
        first = _first_entry(kind, tokens)
        need = _REQUIRED[kind]
        if need not in first:
            self._block("layout_field_missing", f"«{_TOKEN_LABEL[(kind, need)]}» is missing "
                        f"from {section['title']}; the section can't print without it.")
            return None
        title = self.paragraphs[heading].text.strip() or section["title"]
        start = min(s.paragraph_id for s in first.values())

        def opt(name: str) -> OptionalSpan:
            return OptionalSpan(span=first[name]) if name in first else OptionalSpan(present=False)

        def header(names: tuple[str, ...], date_name: str) -> HeaderFieldMapping:
            lead = first.get(names[0]) or min(first.values(), key=lambda s: s.paragraph_id)
            fields = {n: opt(n) for n in names if n in first}
            return HeaderFieldMapping(
                header_paragraph_id=lead.paragraph_id, fields=fields,
                **self._date_alignment(lead, first.get(date_name)),
            )  # fmt: skip

        if kind == "experience":
            return ExperienceMapping(
                heading_paragraph_id=heading, heading_text=title, prototype_entry_start=start,
                header=header(("company", "location", "dates"), "dates"), title=opt("title"),
                title_paragraph_id=first["title"].paragraph_id if "title" in first else None,
                bullet_paragraph_id=first["bullet"].paragraph_id,
            )  # fmt: skip
        if kind == "education":
            return EducationMapping(
                heading_paragraph_id=heading, heading_text=title, prototype_entry_start=start,
                header=header(("school", "location", "dates"), "dates"),
                degree_paragraph_id=first["degree"].paragraph_id,
                detail_paragraph_id=first["detail"].paragraph_id if "detail" in first else None,
            )  # fmt: skip
        if kind == "project":
            return ProjectsMapping(
                heading_paragraph_id=heading, heading_text=title, prototype_entry_start=start,
                header=header(("name", "tech", "date"), "date"), link=opt("link"),
                bullet_paragraph_id=first["bullet"].paragraph_id,
            )  # fmt: skip
        if kind == "skills":
            label = first["label"]
            body = first.get("body")
            if body is None:
                self._block("layout_field_missing", f"«Skills» is missing from {title}.")
                return None
            sep = ": "
            if body.paragraph_id == label.paragraph_id and body.start > label.end:
                sep = self.paragraphs[label.paragraph_id].text[label.end : body.start] or sep
            return SkillsMapping(
                heading_paragraph_id=heading, heading_text=title,
                prototype_paragraph_id=label.paragraph_id, label_span=label, body_span=body,
                separator=sep,
            )  # fmt: skip
        return ListMapping(
            heading_paragraph_id=heading, heading_text=title,
            bullet_paragraph_id=first["item"].paragraph_id,
        )

    def _date_alignment(self, lead: CharSpan, dates: CharSpan | None) -> dict:
        if dates is None:
            return {}
        if dates.paragraph_id != lead.paragraph_id:
            return {"date_alignment": "separate_paragraph", "date_paragraph_id": dates.paragraph_id}
        before = self.paragraphs[dates.paragraph_id].text[: dates.start]
        return {"date_alignment": "tab" if "\t" in before else "inline"}


def _first_entry(kind: str, tokens: list[_Found]) -> dict[str, CharSpan]:
    """Field -> span in the section's first entry. Each field's first occurrence, except
    that a field left only once while the entry lead appears twice belongs to whichever
    entry it sits in — so a field deleted from the first entry is absent, not borrowed
    from the second. (Reading order alone cannot split entries: a table can stack both
    entries' labels in one cell before both bodies in the next.)"""
    def pos(t: _Found) -> tuple[int, int]:
        return (t.span.paragraph_id, t.span.start)

    occurrences: dict[str, list[_Found]] = {}
    for t in sorted(tokens, key=pos):
        occurrences.setdefault(t.token.field, []).append(t)
    leads = occurrences.get(skeleton_tokens.ENTRY_LEAD[kind], [])
    second = pos(leads[1]) if len(leads) > 1 else None
    first: dict[str, CharSpan] = {}
    for name, found in occurrences.items():
        if len(found) == 1 and second is not None and pos(found[0]) > second:
            continue
        first[name] = found[0].span
    return first


def _slug(label: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", label.casefold()).strip("_")[:32] or "line"
    if not base[0].isalpha():
        base = f"line_{base}"[:32]
    key, n = base, 2
    while key in taken:
        key, n = f"{base}_{n}", n + 1
    return key
