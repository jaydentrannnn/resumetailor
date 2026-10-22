"""Read the exported PDF's text back, the way an applicant tracking system will.

Word and LibreOffice can draw a page that looks right while its text layer is empty,
garbled (unmapped glyphs), or interleaved across table columns. This module compares
the text pypdf extracts from ``tailored.pdf`` with the text the docx itself carries:

- the PDF has a text layer, with no unreadable characters;
- the email address and phone number the docx shows are in the PDF text (the phone by
  its digits), so a contact detail only in a link's target is not required;
- every on-page item of 12 or more letters and digits (bullets, titles, degree lines,
  list items: `render.visible_text`) reads back as one contiguous run. Interleaved
  columns or a lost line break inside a word fail this.

Comparison folds case, Unicode compatibility forms (ligatures), soft hyphens, spaces
and punctuation, so wrapping and typography never count as a difference. The must-have
phrases the coverage funnel finds on the page come from these same items, so they need
no check of their own. No model and no document editing: the result only informs the
quality review.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field

from resume_tailor.content.data import MasterResume

#: Items shorter than this (once folded) are skipped: a two-letter skill can match by
#: chance anywhere, and short items are the ones least likely to be split.
MIN_ITEM_CHARS = 12
#: How many missing items a message quotes.
MAX_QUOTED = 2
QUOTE_CHARS = 80
EXTRACTOR = "pypdf"

NO_TEXT = (
    "The PDF has no readable text, so an applicant tracking system would see a blank resume."
)
GARBLED = (
    "Some of the PDF's text reads back as unreadable characters, so an applicant tracking"
    " system may misread it."
)
EMAIL_MISSING = (
    "Your email address isn't in the PDF's text, so an applicant tracking system can't read it."
)
PHONE_MISSING = (
    "Your phone number isn't in the PDF's text, so an applicant tracking system can't read it."
)
UNREADABLE = "The PDF's text couldn't be read back, so it wasn't checked."
NO_PDF = "There is no PDF yet, so its text wasn't checked."


class PdfTextCheck(BaseModel):
    """What the PDF's text layer shows. ``issues`` need review before Fill;
    ``advisories`` are informational."""

    status: Literal["verified", "issues", "unverified"] = "unverified"
    issues: list[str] = Field(default_factory=list)
    advisories: list[str] = Field(default_factory=list)
    extractor: str = EXTRACTOR
    #: The PDF this verdict describes; a re-render writes a new PDF and a new check.
    pdf_sha256: str = ""
    checked_items: int = 0


def fold(text: str) -> str:
    """Letters and digits only, lower-cased, with ligatures and soft hyphens undone."""
    text = unicodedata.normalize("NFKC", text).casefold().replace("\u00ad", "")
    return re.sub(r"[\W_]+", "", text)


def _digits(text: str) -> str:
    return re.sub(r"\D", "", text)


def _quote(text: str) -> str:
    text = " ".join(text.split())
    return text if len(text) <= QUOTE_CHARS else text[: QUOTE_CHARS - 1].rstrip() + "…"


def evaluate(
    pdf_text: str, docx_text: str, *, email: str, phone: str, items: list[str]
) -> PdfTextCheck:
    """The verdict for one PDF's extracted text against its docx's text. Pure."""
    if not fold(pdf_text):
        return PdfTextCheck(status="issues", issues=[NO_TEXT])
    issues: list[str] = []
    if "(cid:" in pdf_text or "\ufffd" in pdf_text:
        issues.append(GARBLED)
    pdf_folded, docx_folded = fold(pdf_text), fold(docx_text)
    address = fold(email)
    if address and address in docx_folded and address not in pdf_folded:
        issues.append(EMAIL_MISSING)
    number = _digits(phone)
    if len(number) >= 7 and number in _digits(docx_text) and number not in _digits(pdf_text):
        issues.append(PHONE_MISSING)
    expected = [
        item for item in dict.fromkeys(items)
        if len(fold(item)) >= MIN_ITEM_CHARS and fold(item) in docx_folded
    ]
    missing = [item for item in expected if fold(item) not in pdf_folded]
    if missing:
        quoted = "; ".join(f"“{_quote(item)}”" for item in missing[:MAX_QUOTED])
        lines = "line" if len(missing) == 1 else "lines"
        issues.append(
            f"{len(missing)} {lines} of your resume don't read back from the PDF as written,"
            f" so an applicant tracking system may garble them: {quoted}"
        )
    return PdfTextCheck(
        status="issues" if issues else "verified", issues=issues, checked_items=len(expected),
    )


def read_pdf_text(pdf_path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages)


def read_docx_text(docx_path: Path) -> str:
    """The docx's visible text: body (tables included), headers and footers, with link
    text."""
    from docx import Document

    from resume_tailor.document import docx_text as text_of

    doc = Document(str(docx_path))
    parts = [text_of.paragraph_text(p) for p, _ in text_of.iter_document_paragraphs(doc)]
    for section in doc.sections:
        for part in (section.header, section.footer):
            parts.extend(text_of.paragraph_text(p) for p in part.paragraphs)
            for table in part.tables:
                parts.extend(
                    text_of.paragraph_text(p)
                    for row in table.rows for cell in row.cells for p in cell.paragraphs
                )
    return "\n".join(parts)


def check(
    pdf_path: Path,
    docx_path: Path,
    *,
    resume: MasterResume,
    bullets: dict[str, str],
    layout: dict | None,
) -> PdfTextCheck:
    """Check the PDF beside a rendered docx. A missing or unreadable file is
    ``unverified``, never a pass."""
    from resume_tailor.document import render

    if not pdf_path.is_file():
        return PdfTextCheck(advisories=[NO_PDF])
    try:
        raw = pdf_path.read_bytes()
        extracted = read_pdf_text(pdf_path)
        source = read_docx_text(docx_path)
    except Exception:  # noqa: BLE001 - any reader failure leaves the PDF unchecked
        return PdfTextCheck(advisories=[UNREADABLE])
    result = evaluate(
        extracted,
        source,
        email=resume.contact.email,
        phone=resume.contact.phone,
        items=render.visible_text(resume, bullets, layout=layout),
    )
    result.pdf_sha256 = hashlib.sha256(raw).hexdigest()
    return result
