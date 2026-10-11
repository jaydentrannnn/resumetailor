"""Your own cover letter: a .docx with «tokens» where the letter's words go.

`starter_letter` gives the person a token version of their current cover letter to start
from; `check` reads an uploaded letter's tokens; `build` turns them into the same Jinja
tags `cover_template.build_cover_template` writes, so `render.render_letter` fills an own
letter exactly like a generated one. `build` is called only from
`cover_template.ensure_cover_template`, which stays the one producer of
`cover_template.docx`.
"""

from __future__ import annotations

import io
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

from ..content.data import MasterResume
from ..pipeline.coverletter_models import CoverLetter
from . import docx_text, render, skeleton_tokens, template_skeleton, template_tagging
from .cover_settings import CoverBlocks, CoverSettings
from .skeleton_tokens import TOKEN_RE, label_text
from .template_profile import CharSpan, TemplateProfile, active_layout
from .template_tags import CUSTOM_FIELD_TAG_FMT
from .template_xml import make_para

#: Letter token label -> (sentinel the starter renders, tag the build writes, block).
LETTER: dict[str, tuple[str, str, str | None]] = {
    "Date": ("QQDATEQQ", "{{ date }}", "date"),
    "Recipient": ("QQRECIPIENTQQ", "{{ line }}", "recipient"),
    "Subject": ("QQSUBJECTQQ", "{{ subject }}", "subject"),
    "Salutation": ("QQSALUTATIONQQ", "{{ salutation }}", "salutation"),
    "Body": ("QQBODYQQ", "{{ para }}", None),
    "Closing": ("QQCLOSINGQQ", "{{ closing }}", "closing"),
    "Signature": ("QQSIGNATUREQQ", "{{ signature }}", "signature"),
}
#: Tokens that repeat their whole line: once per address line, once per paragraph.
_LOOPS = {"Recipient": "{%p for line in inside_address %}", "Body": "{%p for para in paragraphs %}"}
_CONTACT = {t.label: t.field for t in skeleton_tokens.TOKENS if t.kind == "contact"}
_KNOWN = {k.casefold(): k for k in [*LETTER, *_CONTACT, "Name"]}
_COMMENTS = {
    "Date": "Your letter's date. Type «Subject», «Salutation» etc. anywhere; delete a «field» "
    "you don't want. Formatting you give a «field» is kept.",
    "Recipient": "One line of the company's address: this line repeats for each line.",
    "Body": "Your letter's paragraphs: this line repeats once per paragraph.",
}


def contact_key(field_name: str) -> str:
    """The letter context key for one contact field (`render.build_letter_context`)."""
    return f"contact_field_{field_name}"


@dataclass
class _Token:
    paragraph: Paragraph
    pid: int
    start: int
    end: int
    name: str  # "Date", "Email", or "custom:<label>"


def _tokens(doc) -> tuple[list[_Token], list[str]]:
    found: list[_Token] = []
    unknown: list[str] = []
    for pid, (paragraph, _loc) in enumerate(docx_text.iter_document_paragraphs(doc)):
        for m in TOKEN_RE.finditer(paragraph.text):
            name = " ".join((m.group(1) or m.group(2)).split())
            known = _KNOWN.get(name.casefold())
            custom = skeleton_tokens.lookup(name)
            if known is not None:
                found.append(_Token(paragraph, pid, m.start(), m.end(), known))
            elif isinstance(custom, str):
                found.append(_Token(paragraph, pid, m.start(), m.end(), f"custom:{custom}"))
            else:
                unknown.append(name)
    return found, unknown


def check(raw: bytes) -> list[str]:
    """Why `raw` can't be used as a cover letter ([] when it can)."""
    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception as exc:  # noqa: BLE001 - any unreadable package is the same answer
        return [f"The file isn't a readable .docx ({exc})."]
    found, unknown = _tokens(doc)
    problems = [f"«{name}» isn't a field the app knows." for name in unknown]
    names = [t.name for t in found]
    if "Body" not in names:
        problems.append("Add «Body» on its own line where the letter's paragraphs go.")
    for name in LETTER:
        if names.count(name) > 1:
            problems.append(f"«{name}» appears {names.count(name)} times; keep one.")
    for token in found:
        if token.name in _LOOPS and token.paragraph.text.strip() != token.paragraph.text[
            token.start : token.end
        ].strip():
            problems.append(f"«{token.name}» needs a line of its own (it repeats the line).")
    return problems


def _slug(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", label.casefold()).strip("_") or "line"


def build(src: Path, dst: Path, custom_keys: dict[str, str] | None = None) -> None:
    """Tag the own letter at `src` into a cover template at `dst`.

    `custom_keys` maps a custom line's label (casefolded) to its key in the resume, so
    «custom: Street address» prints the person's `street_address` line."""
    custom_keys = custom_keys or {}
    doc = docx.Document(str(src))
    found, _unknown = _tokens(doc)
    by_paragraph: dict[int, list[_Token]] = {}
    for token in found:
        by_paragraph.setdefault(token.pid, []).append(token)
    for pid in sorted(by_paragraph, reverse=True):
        tokens = by_paragraph[pid]
        paragraph = tokens[0].paragraph
        spans = []
        for t in tokens:
            if t.name in LETTER:
                tag = LETTER[t.name][1]
            elif t.name == "Name":
                tag = "{{ name }}"
            elif t.name.startswith("custom:"):
                label = t.name.split(":", 1)[1]
                tag = CUSTOM_FIELD_TAG_FMT % custom_keys.get(label.casefold(), _slug(label))
            else:
                tag = f"{{{{r {contact_key(_CONTACT[t.name])} }}}}"
            spans.append((t.start, t.end, tag, t.name))
        template_tagging.retag_paragraph(paragraph, spans)
        opens, closes = [], []
        for t in tokens:
            block = LETTER.get(t.name, (None, None, None))[2]
            if block is not None:
                opens.append(f'{{%p if blocks["{block}"] %}}')
                closes.insert(0, "{%p endif %}")
            if t.name in _LOOPS:
                opens.append(_LOOPS[t.name])
                closes.insert(0, "{%p endfor %}")
        for text in opens:
            paragraph._p.addprevious(make_para(text))
        for text in reversed(closes):
            paragraph._p.addnext(make_para(text))
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))


def _sentinel_letter() -> CoverLetter:
    return CoverLetter(
        company="Company",
        paragraphs=[LETTER["Body"][0]],
        salutation=LETTER["Salutation"][0],
        closing=LETTER["Closing"][0],
        signature=LETTER["Signature"][0],
        inside_address=[LETTER["Recipient"][0]],
        date=LETTER["Date"][0],
        job_title="Role",
    )


def starter_letter(
    template: Path, profile: TemplateProfile, resume: MasterResume | None
) -> bytes:
    """The cover letter built from the current template, every value a «token»."""
    sentinel = template_skeleton.sentinel_resume(resume, profile)
    settings = CoverSettings(
        blocks=CoverBlocks(subject=True),
        subject_format=LETTER["Subject"][0],
        closing_text=LETTER["Closing"][0],
    )
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "starter.docx"
        render.render_letter(
            sentinel, _sentinel_letter(), template=template, out=out,
            layout=active_layout(profile), settings=settings,
        )  # fmt: skip
        doc = docx.Document(str(out))
    by_sentinel = {sentinel: label for label, (sentinel, _tag, _b) in LETTER.items()}
    for pid, (paragraph, _loc) in enumerate(docx_text.iter_document_paragraphs(doc)):
        text = paragraph.text
        spans = [
            (s, e, "custom", label) if kind == "custom" else (s, e, kind, label)
            for s, e, kind, label in template_skeleton.sentinel_matches(text, header_zone=True)
        ]
        for sentinel_text, label in by_sentinel.items():
            for m in re.finditer(re.escape(sentinel_text), text):
                spans.append((m.start(), m.end(), "letter", label))
        if not spans:
            continue
        template_skeleton._unwrap_hyperlinks(paragraph)
        for start, end, kind, label in sorted(spans, reverse=True):
            if kind == "custom":
                slot = profile.custom_fields[int(label)]
                label = f"custom: {slot.label or slot.key}"
            template_tagging.replace_span_with_tag(
                paragraph, CharSpan(paragraph_id=pid, start=start, end=end), label_text(label)
            )
            if label in _COMMENTS:
                run = next((r for r in paragraph.runs if label_text(label) in r.text), None)
                if run is not None:
                    doc.add_comment(
                        run, text=_COMMENTS[label],
                        author=template_skeleton.COMMENT_AUTHOR, initials="RT",
                    )  # fmt: skip
    _gap_before_salutation(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _gap_before_salutation(doc) -> None:
    """The starter is rendered with the subject line on, which carries the gap under the
    address; give the greeting the same gap so it keeps it when the subject is off."""
    paragraphs = [p for p, _loc in docx_text.iter_document_paragraphs(doc)]
    subject = next((p for p in paragraphs if label_text("Subject") in p.text), None)
    greeting = next((p for p in paragraphs if label_text("Salutation") in p.text), None)
    if subject is not None and greeting is not None and not greeting.paragraph_format.space_before:
        greeting.paragraph_format.space_before = subject.paragraph_format.space_before
