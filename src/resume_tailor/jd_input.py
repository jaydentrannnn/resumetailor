"""Job-description text from a posting URL or an uploaded file (Tailor page, T1).

Both paths return plain text for the student to check and edit before a run: nothing
here calls a model or starts tailoring. Failures raise `JdInputError` carrying a code
and a sentence the UI shows as-is.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath
from urllib.parse import urlparse

from resume_tailor.apply import fetch_jd, identity

#: Same cap as `JobCreateRequest.jd_text`; longer text is cut with a visible warning.
MAX_CHARS = 50_000
#: Below this the "posting" is almost certainly a login wall or an error page.
MIN_CHARS = 200
MAX_FILE_BYTES = 10 * 1024 * 1024
FILE_TYPES = (".txt", ".md", ".pdf", ".docx", ".html", ".htm")

#: Hosts whose postings are only readable when signed in.
_LOGIN_HOSTS = {
    "linkedin.com": "LinkedIn",
    "joinhandshake.com": "Handshake",
}
_CLOSED = re.compile(
    r"no longer (accepting applications|available|open)|this (job|position|posting) "
    r"(has been|is) (closed|filled|expired)|job (not found|has expired)",
    re.I,
)
_ENGLISH = re.compile(r"\b(the|and|with|you|will|for|our|experience)\b", re.I)


class JdInputError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass
class JdText:
    text: str
    source: str
    ats: str = ""
    final_url: str = ""
    warnings: list[str] = field(default_factory=list)


def _login_site(host: str) -> str | None:
    for domain, name in _LOGIN_HOSTS.items():
        if host == domain or host.endswith("." + domain):
            return name
    return None


def _finish(text: str, result: JdText) -> JdText:
    """Normalise, then apply the length checks and warnings both paths share."""
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.splitlines()]
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    if len(text) < MIN_CHARS:
        raise JdInputError(
            "too_short",
            f"Only {len(text)} characters of text were found, which is too little to be a job"
            " posting. Paste the description instead.",
        )
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        result.warnings.append(
            f"The posting was longer than {MAX_CHARS:,} characters, so the end was cut off."
            " Check that the requirements are still there."
        )
    if _CLOSED.search(text[:5000]):
        result.warnings.append(
            "This posting looks closed or expired. You can still tailor to it."
        )
    words = len(text.split())
    if words > 80 and len(_ENGLISH.findall(text)) < words * 0.02:
        result.warnings.append(
            "This posting may not be in English; tailoring works best in English."
        )
    result.text = text
    return result


def from_text(text: str, source: str) -> JdText:
    """Clean text that arrived already extracted (the browser extension's capture)."""
    return _finish(text, JdText(text="", source=source))


def from_url(url: str, *, allow_browser: bool = True) -> JdText:
    url = url.strip()
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise JdInputError("bad_url", "Enter a full web address starting with https://.")
    host = parsed.hostname.lower()
    site = _login_site(host)
    if site:
        raise JdInputError(
            "login_required",
            f"{site} only shows postings to signed-in users, so ResumeTailor can't read this"
            " link. Open the posting, copy the description, and paste it instead.",
        )
    final = identity.resolve_final_url(url)
    fetched = fetch_jd.fetch_jd(
        final, allow_browser=allow_browser, canonical_key=identity.canonical_key(final)
    )
    if fetched.method == "failed":
        if re.search(r"\b(404|410)\b", fetched.error):
            raise JdInputError(
                "closed",
                "That posting no longer exists (the page was not found). It may have closed.",
            )
        raise JdInputError(
            "unreadable",
            "Couldn't read this page. Some job sites only load the description in a browser;"
            " paste the text instead.",
        )
    return _finish(
        fetched.text,
        JdText(text="", source="url", ats=fetched.ats, final_url=fetched.final_url),
    )


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted and not reader.decrypt(""):
            raise JdInputError(
                "encrypted",
                "This PDF is password-protected. Remove the password or paste the text.",
            )
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError) as exc:
        if isinstance(exc, JdInputError):
            raise
        raise JdInputError("unreadable", "This PDF could not be read.") from exc
    if len(text.strip()) < 50:
        raise JdInputError(
            "scanned",
            "This PDF has no text layer (it is a scan or an image). Paste the description instead.",
        )
    return text


def _docx_text(raw: bytes) -> str:
    import zipfile

    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(raw))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise JdInputError("unreadable", "This Word file could not be read.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def from_file(filename: str, raw: bytes) -> JdText:
    suffix = PurePath(filename or "").suffix.lower()
    if suffix not in FILE_TYPES:
        raise JdInputError(
            "file_type", "Upload a .txt, .pdf, .docx or .html file, or paste the text."
        )
    if len(raw) > MAX_FILE_BYTES:
        raise JdInputError("too_large", "That file is over 10 MB; paste the description instead.")
    if suffix == ".pdf":
        text = _pdf_text(raw)
    elif suffix == ".docx":
        text = _docx_text(raw)
    elif suffix in (".html", ".htm"):
        text = fetch_jd.extract_text(_decode(raw))
    else:
        text = _decode(raw)
    return _finish(text, JdText(text="", source=suffix.lstrip(".")))
