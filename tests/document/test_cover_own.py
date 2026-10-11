"""Your own cover letter: starter download, checks, build and render."""

from __future__ import annotations

import io
import zipfile

import docx
import pytest

from resume_tailor import config
from resume_tailor.content.data import CustomField
from resume_tailor.document import (
    cover_own,
    cover_template,
    docx_text,
    render,
    template_analyze,
    template_profile,
)
from resume_tailor.document.cover_settings import CoverBlocks, CoverSettings
from resume_tailor.pipeline.coverletter_models import CoverLetter
from tests.fixtures import _docx_bytes, _full_featured_resume, _table_resume, synthetic_resume


@pytest.fixture
def setup(tmp_path, monkeypatch):
    """(generated cover template, profile) for a fixture builder."""

    def make(builder):
        raw = _docx_bytes(builder)
        profile = template_analyze.analyze_docx(raw=raw).suggested_profile
        src, dst = tmp_path / "baseline.docx", tmp_path / "cover.docx"
        src.write_bytes(raw)
        monkeypatch.setattr(config, "BASELINE_TEMPLATE_PATH", src)
        cover_template.build_cover_template(src, dst, profile=profile)
        return dst, profile

    return make


def _texts(raw: bytes) -> list[str]:
    doc = docx.Document(io.BytesIO(raw))
    return [p.text for p, _loc in docx_text.iter_document_paragraphs(doc)]


def _letter(resume) -> CoverLetter:
    return CoverLetter(
        company="Example Corp",
        paragraphs=["First paragraph.", "Second paragraph."],
        salutation="Dear Hiring Manager,",
        closing="Sincerely,",
        signature=resume.contact.name,
        inside_address=["Example Corp", "Irvine, CA"],
        date="August 29, 2026",
        job_title="Data Analyst",
    )


def _render_own(tmp_path, raw: bytes, profile, resume=None, settings=None, keys=None):
    own, tpl, out = tmp_path / "own.docx", tmp_path / "own_tpl.docx", tmp_path / "out.docx"
    own.write_bytes(raw)
    cover_own.build(own, tpl, custom_keys=keys)
    resume = resume or synthetic_resume()
    render.render_letter(
        resume, _letter(resume), template=tpl, out=out,
        layout=template_profile.active_layout(profile), settings=settings or CoverSettings(),
    )  # fmt: skip
    return [p.text for p, _loc in docx_text.iter_document_paragraphs(docx.Document(str(out)))]


def test_starter_letter_has_a_field_for_every_value_and_notes(setup):
    tpl, profile = setup(_full_featured_resume)
    raw = cover_own.starter_letter(tpl, profile, synthetic_resume())
    texts = _texts(raw)
    for token in ("«Name»", "«Date»", "«Recipient»", "«Subject»", "«Salutation»", "«Body»",
                  "«Closing»", "«Signature»"):  # fmt: skip
        assert token in texts, token
    assert any("«Email»" in t for t in texts)
    xml = zipfile.ZipFile(io.BytesIO(raw)).read("word/document.xml").decode()
    assert "QQ" not in xml and "{{" not in xml and "{%" not in xml
    assert "word/comments.xml" in zipfile.ZipFile(io.BytesIO(raw)).namelist()
    assert cover_own.check(raw) == []


def test_starter_round_trip_renders_the_letter(setup, tmp_path):
    tpl, profile = setup(_full_featured_resume)
    raw = cover_own.starter_letter(tpl, profile, synthetic_resume())
    resume = synthetic_resume()
    texts = _render_own(tmp_path, raw, profile, settings=CoverSettings())
    for line in ("August 29, 2026", "Example Corp", "Irvine, CA", "Dear Hiring Manager,",
                 "First paragraph.", "Second paragraph.", "Sincerely,", resume.contact.name):  # fmt: skip
        assert line in texts, line
    assert any(resume.contact.email in t for t in texts)
    assert not any("«" in t or "{{" in t for t in texts)


def test_formatting_given_to_a_field_is_kept(setup, tmp_path):
    tpl, profile = setup(_full_featured_resume)
    doc = docx.Document(io.BytesIO(cover_own.starter_letter(tpl, profile, synthetic_resume())))
    for paragraph, _loc in docx_text.iter_document_paragraphs(doc):
        for run in paragraph.runs:
            if run.text == "«Date»":
                run.bold = True
    buffer = io.BytesIO()
    doc.save(buffer)
    own, out_tpl, out = tmp_path / "o.docx", tmp_path / "t.docx", tmp_path / "l.docx"
    own.write_bytes(buffer.getvalue())
    cover_own.build(own, out_tpl)
    resume = synthetic_resume()
    render.render_letter(resume, _letter(resume), template=out_tpl, out=out,
                         layout=template_profile.active_layout(profile))  # fmt: skip
    date = next(
        p for p, _loc in docx_text.iter_document_paragraphs(docx.Document(str(out)))
        if p.text == "August 29, 2026"
    )
    assert date.runs[0].bold


def test_greeting_keeps_its_gap_under_the_address_without_a_subject(setup, tmp_path):
    tpl, profile = setup(_full_featured_resume)
    own, out_tpl, out = tmp_path / "o.docx", tmp_path / "t.docx", tmp_path / "l.docx"
    own.write_bytes(cover_own.starter_letter(tpl, profile, synthetic_resume()))
    cover_own.build(own, out_tpl)
    resume = synthetic_resume()
    render.render_letter(resume, _letter(resume), template=out_tpl, out=out,
                         layout=template_profile.active_layout(profile))  # fmt: skip
    greeting = next(
        p for p, _loc in docx_text.iter_document_paragraphs(docx.Document(str(out)))
        if p.text == "Dear Hiring Manager,"
    )
    assert greeting.paragraph_format.space_before


def test_a_block_turned_off_removes_its_line(setup, tmp_path):
    tpl, profile = setup(_full_featured_resume)
    raw = cover_own.starter_letter(tpl, profile, synthetic_resume())
    settings = CoverSettings(blocks=CoverBlocks(date=False, recipient=False))
    texts = _render_own(tmp_path, raw, profile, settings=settings)
    assert "August 29, 2026" not in texts and "Irvine, CA" not in texts
    assert "First paragraph." in texts


def test_table_letterhead_custom_line_reaches_the_letter(setup, tmp_path):
    tpl, profile = setup(_table_resume)
    raw = cover_own.starter_letter(tpl, profile, synthetic_resume())
    assert "«custom: Street address»" in _texts(raw)
    resume = synthetic_resume().model_copy(
        update={"custom_fields": [CustomField(key="street_address", value="9 New Rd")]}
    )
    keys = {(s.label or s.key).casefold(): s.key for s in profile.custom_fields}
    assert "9 New Rd" in _render_own(tmp_path, raw, profile, resume=resume, keys=keys)


def _letter_doc(*lines: str) -> bytes:
    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "lines, problem",
    [
        (["«Date»", "Dear team"], "Add «Body»"),
        (["Hello «Body»"], "line of its own"),
        (["«Date»", "«Body»", "<<date>>"], "appears 2 times"),
        (["«Body»", "«Fax»"], "«Fax»"),
    ],
)
def test_check_explains_what_is_wrong(lines, problem):
    problems = cover_own.check(_letter_doc(*lines))
    assert any(problem in p for p in problems), problems


def test_a_hand_made_letter_with_typed_fields_works(tmp_path, setup):
    _tpl, profile = setup(_full_featured_resume)
    raw = _letter_doc("<<Name>> — <<Email>>", "<<date>>", "<<Salutation>>", "<<Body>>",
                      "Best,", "<<Signature>>")  # fmt: skip
    assert cover_own.check(raw) == []
    resume = synthetic_resume()
    texts = _render_own(tmp_path, raw, profile)
    assert f"{resume.contact.name} — {resume.contact.email}" in texts
    assert texts.count("First paragraph.") == 1 and "Best," in texts


def test_ensure_switches_between_own_and_generated(setup, tmp_path, monkeypatch):
    tpl, profile = setup(_full_featured_resume)
    own = tmp_path / "cover_own.docx"
    own.write_bytes(_letter_doc("MY OWN LETTER", "«Body»"))
    monkeypatch.setattr(config, "COVER_OWN_PATH", own)
    dst, meta = tmp_path / "live.docx", tmp_path / "live.meta.json"
    kwargs = dict(src=config.BASELINE_TEMPLATE_PATH, dst=dst, meta_path=meta, profile=profile)
    cover_template.ensure_cover_template(**kwargs, settings=CoverSettings(design="own"))
    assert "MY OWN LETTER" in _texts(dst.read_bytes())
    cover_template.ensure_cover_template(**kwargs, settings=CoverSettings())
    assert "MY OWN LETTER" not in _texts(dst.read_bytes())
    own.unlink()  # "own" without a file falls back to the generated letter
    cover_template.ensure_cover_template(**kwargs, settings=CoverSettings(design="own"))
    assert "{{ date }}" in _texts(dst.read_bytes())
