"""Bullet markers keep the glyph the upload shows: Symbol/Wingdings code points become
their Unicode twin before the marker font is pinned, and a paragraph is only moved onto
the canonical list when that list draws the same marker."""

from __future__ import annotations

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from resume_tailor.document import template_bullets, template_tags, template_xml

_SYMBOL_BULLET = chr(0xF0B7)
_BULLET = chr(0x2022)


def _add_list(document, num_id: str, glyph: str, font: str | None = None) -> str:
    """A bullet list instance `num_id` whose lvl0 marker is `glyph` (in `font`)."""
    root = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), f"a{num_id}")
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), glyph)
    lvl.extend([num_fmt, text])
    if font:
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), font)
        r_pr.append(r_fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    first_num = root.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), num_id)
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), f"a{num_id}")
    num.append(abs_ref)
    root.append(num)
    return num_id


def _lvl0(document, num_id: str):
    root = document.part.numbering_part.element
    for anum in root.findall(qn("w:abstractNum")):
        if anum.get(qn("w:abstractNumId")) == f"a{num_id}":
            return anum.find(qn("w:lvl"))
    raise AssertionError(num_id)


def _glyph_and_font(document, num_id: str) -> tuple[str, str | None]:
    lvl = _lvl0(document, num_id)
    r_pr = lvl.find(qn("w:rPr"))
    r_fonts = r_pr.find(qn("w:rFonts")) if r_pr is not None else None
    font = r_fonts.get(qn("w:ascii")) if r_fonts is not None else None
    return lvl.find(qn("w:lvlText")).get(qn("w:val")), font


def _document():
    document = docx.Document()
    document.add_paragraph("- placeholder", style="List Bullet")  # forces a numbering part
    return document


def test_symbol_bullet_becomes_a_unicode_bullet_in_the_pinned_font():
    document = _document()
    _add_list(document, "71", _SYMBOL_BULLET, "Symbol")
    _add_list(document, "72", "-", "Lora")

    template_bullets.normalize_bullet_numbering(document)

    assert _glyph_and_font(document, "71") == (_BULLET, template_tags.NOTO_MARKER_FONT)
    assert _glyph_and_font(document, "72") == ("-", template_tags.NOTO_MARKER_FONT)


def test_unknown_symbol_font_marker_keeps_its_own_font():
    document = _document()
    _add_list(document, "73", chr(0xF06C), "Wingdings")

    template_bullets.normalize_bullet_numbering(document)

    assert _glyph_and_font(document, "73") == (chr(0xF06C), "Wingdings")


def _bullet_paragraph(document, num_id: str):
    paragraph = document.add_paragraph("Did a thing.")
    template_xml.set_num_id(paragraph, num_id)
    return paragraph


def _num_id(paragraph) -> str:
    return paragraph._p.pPr.numPr.numId.get(qn("w:val"))


def test_retarget_keeps_a_section_whose_marker_differs():
    document = _document()
    _add_list(document, "74", "-", template_tags.NOTO_MARKER_FONT)
    _add_list(document, "75", _SYMBOL_BULLET, "Symbol")
    paragraph = _bullet_paragraph(document, "75")

    template_bullets.retarget_bullet(document, paragraph, "74")

    assert _num_id(paragraph) == "75"


def test_retarget_merges_lists_that_draw_the_same_marker():
    document = _document()
    _add_list(document, "76", _BULLET, template_tags.NOTO_MARKER_FONT)
    _add_list(document, "77", _SYMBOL_BULLET, "Symbol")  # the same "•" on the page
    paragraph = _bullet_paragraph(document, "77")

    template_bullets.retarget_bullet(document, paragraph, "76")

    assert _num_id(paragraph) == "76"


def test_retarget_gives_a_plain_paragraph_the_canonical_list():
    document = _document()
    _add_list(document, "78", "-", template_tags.NOTO_MARKER_FONT)
    paragraph = document.add_paragraph("Bachelor of Science")

    template_bullets.retarget_bullet(document, paragraph, "78")

    assert _num_id(paragraph) == "78"
