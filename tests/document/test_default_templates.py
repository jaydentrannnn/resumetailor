"""Built-in starter templates (`default_templates`, plan P3-T)."""

from __future__ import annotations

import io

import docx
import pytest

from resume_tailor.document import default_templates, render, template_analyze, template_profile
from tests.fixtures import synthetic_resume


@pytest.mark.parametrize("name", default_templates.names())
def test_each_design_builds_reproducibly_and_analyzes_cleanly(name):
    raw = default_templates.build(name)
    assert raw == default_templates.build(name)
    result = template_analyze.analyze_docx(raw=raw)
    assert result.issues == []
    profile = result.suggested_profile
    assert profile is not None
    # Generic mode: any number of sections, in the student's order, with their titles.
    assert profile.section_mode == "generic"
    assert {s.kind for s in profile.sections} == {
        "education",
        "experience",
        "project",
        "skills",
        "list",
    }


@pytest.mark.parametrize("name", default_templates.names())
def test_each_design_tags_and_renders_a_resume(name, tmp_path):
    written = default_templates.write_bundle(name, tmp_path / name)
    baseline, profile_file, tagged = written
    assert [p.name for p in written] == [
        "original_export.docx",
        "template_profile.json",
        "main_template.docx",
    ]
    assert baseline.read_bytes() == default_templates.build(name)
    profile = template_profile.load_profile(profile_file)
    out = render.render(
        synthetic_resume(),
        template=tagged,
        out=tmp_path / "filled.docx",
        layout=template_profile.active_layout(profile),
    )
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert synthetic_resume().contact.name in text
    # None of the template's sample content leaks into a student's resume.
    assert "Northwind Capital" not in text
    assert "Alex Doe" not in text


def test_sample_content_round_trips_through_resume_import():
    resume = default_templates.sample_resume("business")
    assert resume.contact.name == "Alex Doe"
    titles = [s.title for s in resume.sections]
    assert titles[0] == "EDUCATION"
    assert "LEADERSHIP & ACTIVITIES" in titles


def test_business_puts_education_first_and_caps_the_name_by_formatting():
    spec = default_templates.design("business")
    assert spec.education_first
    assert spec.sections[0] == "education"
    document = docx.Document(io.BytesIO(default_templates.build("business")))
    name_run = document.paragraphs[0].runs[0]
    assert name_run.text == "Alex Doe"
    assert name_run.font.all_caps is True


def test_unknown_name_raises():
    with pytest.raises(default_templates.UnknownTemplate):
        default_templates.build("fancy")
