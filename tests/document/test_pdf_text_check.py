"""The exported PDF's text read back: contact details, garbling, split lines."""

from pathlib import Path

from docx import Document

from resume_tailor.document import pdf_text_check, render
from resume_tailor.pipeline import resume_quality
from tests.fixtures import synthetic_resume
from tests.pdf_fixtures import T, make_pdf

BULLET = "Built a caching layer that cut page load time by 40 percent"
DOCX = f"Jordan Rivera\njordan@example.com | (555) 123-4567\nExperience\n{BULLET}"


def _evaluate(pdf_text: str, *, docx: str = DOCX, items=(BULLET,)):
    return pdf_text_check.evaluate(
        pdf_text, docx, email="jordan@example.com", phone="(555) 123-4567", items=list(items)
    )


def test_matching_text_verifies_despite_wraps_ligatures_and_soft_hyphens():
    pdf = (
        "Jordan Rivera\njordan@example.com • 555.123.4567\nEXPERIENCE\n"
        "• Built a caching layer that cut page load time by 40 per\u00adcent"
    ).replace("load time", "load\ntime")
    found = _evaluate(pdf)
    assert found.status == "verified"
    assert found.issues == [] and found.checked_items == 1
    ligature = pdf.replace("caching layer", "e\ufb03cient caching layer")
    docx = DOCX.replace("caching layer", "efficient caching layer")
    item = BULLET.replace("caching layer", "efficient caching layer")
    assert _evaluate(ligature, docx=docx, items=[item]).status == "verified"


def test_no_text_layer_is_one_issue():
    found = _evaluate("  \n ")
    assert found.status == "issues"
    assert found.issues == [pdf_text_check.NO_TEXT]


def test_unmapped_glyphs_are_an_issue():
    found = _evaluate(f"(cid:12)(cid:7) {DOCX}")
    assert pdf_text_check.GARBLED in found.issues
    assert pdf_text_check.GARBLED in _evaluate(DOCX + " \ufffd").issues


def test_contact_details_shown_in_the_docx_must_reach_the_pdf_text():
    found = _evaluate(f"Jordan Rivera\nExperience\n{BULLET}")
    assert found.issues == [pdf_text_check.EMAIL_MISSING, pdf_text_check.PHONE_MISSING]


def test_contact_details_the_docx_never_shows_are_not_required():
    # An email only behind an "Email" link's target is not text on the page.
    docx = f"Jordan Rivera\nEmail | LinkedIn\n{BULLET}"
    assert _evaluate(f"Jordan Rivera\nEmail LinkedIn\n{BULLET}", docx=docx).status == "verified"


def test_interleaved_columns_name_the_split_lines():
    other = "Mentored four interns through their first production launch"
    docx = f"{DOCX}\n{other}"
    pdf = (
        "jordan@example.com (555) 123-4567\n"
        "Built a caching layer that cut Mentored four interns through\n"
        "page load time by 40 percent their first production launch"
    )
    found = _evaluate(pdf, docx=docx, items=[BULLET, other, "SQL"])
    assert found.status == "issues" and found.checked_items == 2
    [message] = found.issues
    assert message.startswith("2 lines of your resume don't read back")
    assert f"“{BULLET}”" in message


def test_items_the_docx_does_not_carry_are_skipped():
    found = _evaluate(DOCX, items=[BULLET, "A bullet this template never prints"])
    assert found.status == "verified" and found.checked_items == 1


def _docx(path: Path, lines: list[str], *, header: str = "") -> Path:
    doc = Document()
    if header:
        doc.sections[0].header.paragraphs[0].text = header
    for line in lines:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Skills table cell"
    doc.save(path)
    return path


def _resume_run(tmp_path: Path, *, pdf_lines: list[str]):
    resume = synthetic_resume()
    bullets = {b.id: b.text for b in resume.all_bullets()}
    items = render.visible_text(resume, bullets, layout={})
    docx = _docx(tmp_path / "tailored.docx", items, header="jordan@example.com (555) 123-4567")
    pdf = tmp_path / "tailored.pdf"
    pdf.write_bytes(make_pdf([[T(40, 40 + 12 * i, line) for i, line in enumerate(pdf_lines)]]))
    return resume, bullets, docx, pdf, items


def test_check_reads_both_files_and_binds_the_pdf_hash(tmp_path):
    resume, bullets, docx, pdf, items = _resume_run(tmp_path, pdf_lines=[])
    pdf.write_bytes(make_pdf([[T(40, 40, "jordan@example.com (555) 123-4567")]
                              + [T(40, 60 + 12 * i, line) for i, line in enumerate(items)]]))
    found = pdf_text_check.check(pdf, docx, resume=resume, bullets=bullets, layout={})
    assert found.status == "verified", found.issues
    assert found.checked_items > 3
    import hashlib

    assert found.pdf_sha256 == hashlib.sha256(pdf.read_bytes()).hexdigest()
    assert "jordan@example.com" in pdf_text_check.read_docx_text(docx)
    assert "Skills table cell" in pdf_text_check.read_docx_text(docx)


def test_check_flags_an_email_only_the_docx_has(tmp_path):
    resume, bullets, docx, pdf, items = _resume_run(tmp_path, pdf_lines=[])
    pdf.write_bytes(make_pdf([[T(40, 40, "(555) 123-4567")]
                              + [T(40, 60 + 12 * i, line) for i, line in enumerate(items)]]))
    found = pdf_text_check.check(pdf, docx, resume=resume, bullets=bullets, layout={})
    assert found.issues == [pdf_text_check.EMAIL_MISSING]


def test_missing_or_unreadable_pdf_is_unverified_never_a_pass(tmp_path):
    resume, bullets, docx, pdf, _ = _resume_run(tmp_path, pdf_lines=["x"])
    missing = pdf_text_check.check(
        tmp_path / "none.pdf", docx, resume=resume, bullets=bullets, layout={}
    )
    assert (missing.status, missing.advisories) == ("unverified", [pdf_text_check.NO_PDF])
    pdf.write_bytes(b"not a pdf")
    broken = pdf_text_check.check(pdf, docx, resume=resume, bullets=bullets, layout={})
    assert (broken.status, broken.advisories) == ("unverified", [pdf_text_check.UNREADABLE])


def test_quality_warnings_carry_pdf_issues_but_not_advisories():
    quality = resume_quality.ResumeQuality(
        pdf_text=pdf_text_check.PdfTextCheck(
            status="issues", issues=[pdf_text_check.EMAIL_MISSING], advisories=["note"]
        )
    )
    assert quality.warnings == [pdf_text_check.EMAIL_MISSING]
    unverified = resume_quality.ResumeQuality(
        pdf_text=pdf_text_check.PdfTextCheck(advisories=[pdf_text_check.UNREADABLE])
    )
    assert unverified.warnings == []
    legacy = resume_quality.ResumeQuality.model_validate_json('{"verified": true}')
    assert legacy.pdf_text is None and legacy.warnings == []
