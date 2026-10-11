"""Unmapped lines above the first section, kept as editable custom fields."""

from __future__ import annotations

import zipfile
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn
from pydantic import ValidationError

from resume_tailor.content.data import CustomField, MasterResume
from resume_tailor.document import (
    cover_template,
    profile_validation,
    render,
    template_analyze,
    template_build,
    template_profile,
    template_tags,
)
from resume_tailor.document.template_profile import TemplateProfile
from resume_tailor.pipeline.coverletter_models import CoverLetter
from tests.fixtures import (
    _add_bullet_numbering,
    _add_hyperlink,
    _docx_bytes,
    _make_bullet,
    _table_resume,
    synthetic_resume,
)


def _tagline_resume(document) -> None:
    """Paragraph layout with two extra lines between the contact line and EDUCATION."""
    num_id = _add_bullet_numbering(document)
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("London • ada@example.com")
    document.add_paragraph("Engineer who ships numerical tools")
    linked = document.add_paragraph("Portfolio: ")
    _add_hyperlink(linked, "ada.dev", "https://ada.dev")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("University of London | UK\t2018 - 2022")
    _make_bullet(document, "BSc Computer Science", num_id)
    document.add_paragraph("WORK EXPERIENCES")
    document.add_paragraph("Analytical Engines | London\t2022 - Present")
    document.add_paragraph("Software Engineer")
    _make_bullet(document, "Built numerical engines in Python.", num_id)


def _build(tmp_path: Path, build, profile_update: dict | None = None):
    raw = _docx_bytes(build)
    analysis = template_analyze.analyze_docx(raw=raw)
    profile = analysis.suggested_profile
    assert profile is not None, analysis.issues
    if profile_update:
        profile = profile.model_copy(update=profile_update)
    src, dst = tmp_path / "baseline.docx", tmp_path / "main_template.docx"
    src.write_bytes(raw)
    template_build.build_from_profile(src, dst, profile)
    return analysis, profile, src, dst


def _render_texts(resume: MasterResume, template: Path, profile: TemplateProfile, out: Path):
    render.render(resume, template=template, out=out, layout=template_profile.active_layout(profile))
    doc = docx.Document(str(out))
    body = [p.text for p in doc.paragraphs]
    cells = [p.text for t in doc.tables for c in t._cells for p in c.paragraphs]
    return body + cells


# -- analysis --------------------------------------------------------------------------


def test_table_contact_leftover_is_suggested_as_a_custom_field():
    analysis = template_analyze.analyze_docx(raw=_docx_bytes(_table_resume))
    assert [(u.text, u.suggestion, u.suggested_key) for u in analysis.unmapped] == [
        ("123 Example St.", "field", "street_address")
    ]
    [slot] = analysis.suggested_profile.custom_fields
    assert (slot.key, slot.label, slot.text) == ("street_address", "Street address", "123 Example St.")


def test_paragraph_layout_extra_lines_are_listed_and_a_linked_one_stays_static():
    analysis = template_analyze.analyze_docx(raw=_docx_bytes(_tagline_resume))
    found = {u.text: u.suggestion for u in analysis.unmapped}
    assert found == {"Engineer who ships numerical tools": "field", "Portfolio: ada.dev": "static"}
    assert [f.text for f in analysis.suggested_profile.custom_fields] == [
        "Engineer who ships numerical tools"
    ]


# -- build + render --------------------------------------------------------------------


def test_custom_line_renders_template_text_until_the_resume_overrides_it(tmp_path):
    _a, profile, _src, dst = _build(tmp_path, _table_resume)
    resume = synthetic_resume()
    assert "123 Example St." in _render_texts(resume, dst, profile, tmp_path / "a.docx")

    edited = resume.model_copy(
        update={"custom_fields": [CustomField(key="street_address", value="9 New Rd")]}
    )
    texts = _render_texts(edited, dst, profile, tmp_path / "b.docx")
    assert "9 New Rd" in texts and "123 Example St." not in texts


def test_clearing_a_body_line_removes_the_paragraph(tmp_path):
    _a, profile, _src, dst = _build(tmp_path, _tagline_resume)
    [slot] = profile.custom_fields
    tagged = [p.text for p in docx.Document(str(dst)).paragraphs]
    assert template_tags.CUSTOM_FIELD_IF_FMT % slot.key in tagged

    resume = synthetic_resume()
    before = _render_texts(resume, dst, profile, tmp_path / "a.docx")
    cleared = resume.model_copy(update={"custom_fields": [CustomField(key=slot.key, value="")]})
    after = _render_texts(cleared, dst, profile, tmp_path / "b.docx")
    assert "Engineer who ships numerical tools" in before
    assert len(after) == len(before) - 1


def test_clearing_a_table_cell_line_keeps_the_cell_valid(tmp_path):
    _a, profile, _src, dst = _build(tmp_path, _table_resume)
    resume = synthetic_resume().model_copy(
        update={"custom_fields": [CustomField(key="street_address", value="")]}
    )
    out = tmp_path / "r.docx"
    render.render(resume, template=dst, out=out, layout=template_profile.active_layout(profile))
    for table in docx.Document(str(out)).tables:
        for row in table.rows:
            for tc in row._tr.findall(qn("w:tc")):
                assert tc.find(qn("w:p")) is not None


def test_dropped_line_is_removed_and_kept_line_stays_literal(tmp_path):
    analysis = template_analyze.analyze_docx(raw=_docx_bytes(_tagline_resume))
    by_text = {u.text: u.paragraph_id for u in analysis.unmapped}
    _a, profile, _src, dst = _build(
        tmp_path,
        _tagline_resume,
        {"custom_fields": [], "dropped_paragraph_ids": [by_text["Portfolio: ada.dev"]]},
    )
    texts = [p.text for p in docx.Document(str(dst)).paragraphs]
    assert "Engineer who ships numerical tools" in texts  # "keep as-is": untouched
    assert not any("Portfolio" in t for t in texts)


def test_profile_without_custom_fields_leaves_the_line_untouched(tmp_path):
    _a, _profile, _src, dst = _build(tmp_path, _table_resume, {"custom_fields": []})
    xml = zipfile.ZipFile(dst).read("word/document.xml").decode("utf-8")
    assert "123 Example St." in xml and "custom[" not in xml


# -- validation --------------------------------------------------------------------------


def test_custom_field_on_the_contact_line_is_rejected():
    analysis = template_analyze.analyze_docx(raw=_docx_bytes(_table_resume))
    profile = analysis.suggested_profile
    bad = profile.model_copy(
        update={
            "custom_fields": [
                profile.custom_fields[0].model_copy(update={"paragraph_id": profile.contact.paragraph_id})
            ]
        }
    )
    paras = {p.id: p for p in analysis.paragraphs}

    def codes(prof):
        return {i.code for i in profile_validation._ProfileValidator(prof, paras).run()}

    assert "bad_custom_line" not in codes(profile)
    assert "bad_custom_line" in codes(bad)


def test_profile_rejects_a_paragraph_both_kept_and_dropped():
    profile = template_analyze.analyze_docx(raw=_docx_bytes(_table_resume)).suggested_profile
    with pytest.raises(ValidationError):
        TemplateProfile.model_validate(
            {
                **profile.model_dump(),
                "dropped_paragraph_ids": [profile.custom_fields[0].paragraph_id],
            }
        )


def test_resume_rejects_duplicate_custom_keys():
    data = synthetic_resume().model_dump(by_alias=True)
    data["custom_fields"] = [{"key": "a", "value": "1"}, {"key": "a", "value": "2"}]
    with pytest.raises(ValidationError):
        MasterResume.model_validate(data)


# -- cover letter ------------------------------------------------------------------------


def test_cover_letterhead_carries_custom_line_and_contact_slots(tmp_path):
    raw = _docx_bytes(_table_resume)
    profile = template_analyze.analyze_docx(raw=raw).suggested_profile
    src, dst = tmp_path / "baseline.docx", tmp_path / "cover.docx"
    src.write_bytes(raw)
    cover_template.build_cover_template(src, dst, profile=profile)

    resume = synthetic_resume().model_copy(
        update={"custom_fields": [CustomField(key="street_address", value="9 New Rd")]}
    )
    letter = CoverLetter(
        company="Example Corp", company_location="", addressee="", paragraphs=["Body."],
        salutation="Dear Hiring Manager,", closing="Sincerely,",
        signature=resume.contact.name, inside_address=["Example Corp"], date="May 1, 2026",
    )
    out = tmp_path / "letter.docx"
    render.render_letter(
        resume, letter, template=dst, out=out, layout=template_profile.active_layout(profile)
    )
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "9 New Rd" in xml and "123 Example St." not in xml
    # The table letterhead's slot tags used to render empty: the letter context never
    # set `contact_slot_<i>`.
    assert resume.contact.email in xml
