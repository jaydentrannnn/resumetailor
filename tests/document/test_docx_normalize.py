"""Upload clean-up (`docx_normalize`), header contact, and content-only import (P3-D)."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from resume_tailor.document import (
    analysis_types,
    convert,
    cover_template,
    docx_normalize,
    template_analyze,
    template_build,
    template_verify,
)
from resume_tailor.importing import import_layout, resume_import
from tests.fixtures import (
    _add_bullet_numbering,
    _docx_bytes,
    _make_bullet,
    _sidebar_table_resume,
    _standard_resume,
)


def _body_texts(raw: bytes) -> list[str]:
    return [p.text for p in docx.Document(io.BytesIO(raw)).paragraphs]


def _codes(prepared) -> set[str]:
    return {n.code for n in prepared.notices}


def test_a_clean_file_comes_back_byte_for_byte():
    raw = _docx_bytes(_standard_resume)
    prepared = docx_normalize.prepare(raw, "resume.docx")
    assert prepared.raw == raw
    assert prepared.notices == []


def _with_tracked_changes(document) -> None:
    _standard_resume(document)
    target = next(p for p in document.paragraphs if p.text == "Software Engineer")
    target._p.append(
        parse_xml(
            f'<w:ins {nsdecls("w")} w:id="1" w:author="x"><w:r><w:t xml:space="preserve">'
            " II</w:t></w:r></w:ins>"
        )
    )
    target._p.append(
        parse_xml(
            f'<w:del {nsdecls("w")} w:id="2" w:author="x"><w:r><w:delText> (intern)'
            "</w:delText></w:r></w:del>"
        )
    )


def test_tracked_changes_are_accepted_deterministically():
    raw = _docx_bytes(_with_tracked_changes)
    first = docx_normalize.prepare(raw, "resume.docx")
    second = docx_normalize.prepare(raw, "resume.docx")
    assert first.raw == second.raw
    assert "Software Engineer II" in _body_texts(first.raw)
    assert "tracked_changes_accepted" in _codes(first)
    xml = docx.Document(io.BytesIO(first.raw)).element.xml
    assert "w:ins " not in xml and "w:del " not in xml and "delText" not in xml


def _with_comment_and_control(document) -> None:
    _standard_resume(document)
    target = next(p for p in document.paragraphs if p.text == "Software Engineer")
    target._p.insert(1, parse_xml(f'<w:commentRangeStart {nsdecls("w")} w:id="0"/>'))
    target._p.append(parse_xml(f'<w:commentRangeEnd {nsdecls("w")} w:id="0"/>'))
    target._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:commentReference w:id="0"/></w:r>'))
    # A body-level content control around the PROJECTS heading.
    heading = next(p for p in document.paragraphs if p.text == "PROJECTS")
    sdt = parse_xml(f"<w:sdt {nsdecls('w')}><w:sdtPr/><w:sdtContent/></w:sdt>")
    heading._p.addprevious(sdt)
    sdt.find(qn("w:sdtContent")).append(heading._p)


def test_comments_and_content_controls_are_removed():
    raw = _docx_bytes(_with_comment_and_control)
    # Before: the analyzer can't see the heading inside the content control.
    before = template_analyze.analyze_docx(raw=raw)
    assert "PROJECTS" not in [s.heading_text for s in before.sections]

    prepared = docx_normalize.prepare(raw, "resume.docx")
    assert {"comments_removed", "content_controls_unwrapped"} <= _codes(prepared)
    after = template_analyze.analyze_docx(raw=prepared.raw)
    assert "PROJECTS" in [s.heading_text for s in after.sections]
    xml = docx.Document(io.BytesIO(prepared.raw)).element.xml
    assert "commentReference" not in xml and "w:sdt>" not in xml


def _typed_bullets(document) -> None:
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("London • ada@example.com")
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Analytical Engines | London\t2022 - Present")
    document.add_paragraph("Software Engineer")
    document.add_paragraph("•\tBuilt numerical engines in Python.")
    document.add_paragraph("• Wrote the first program.")
    document.add_paragraph("-5% is not a bullet")


def test_typed_bullets_become_a_real_list_only_when_asked():
    raw = _docx_bytes(_typed_bullets)
    assert docx_normalize.prepare(raw, "resume.docx").raw == raw

    prepared = docx_normalize.prepare(raw, "resume.docx", convert_bullets=True)
    assert "typed_bullets_converted" in _codes(prepared)
    document = docx.Document(io.BytesIO(prepared.raw))
    bullets = [p for p in document.paragraphs if analysis_types.is_bullet(p)]
    assert [p.text for p in bullets] == [
        "Built numerical engines in Python.",
        "Wrote the first program.",
    ]
    assert any(p.text == "-5% is not a bullet" for p in document.paragraphs)
    # Only the "-5%" line (paragraph 7) still reads as a typed bullet to the
    # analyzer's loose check; the converted lines are real list items now.
    analysis = template_analyze.analyze_docx(raw=prepared.raw)
    flagged = [i.message for i in analysis.issues if i.code == "manual_bullets"]
    assert all(m.startswith("Paragraph 7 ") for m in flagged)


def test_typed_bullets_in_a_file_without_a_numbering_part():
    src = zipfile.ZipFile(io.BytesIO(_docx_bytes(_typed_bullets)))
    assert "word/numbering.xml" in src.namelist()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            if info.filename == "word/numbering.xml":
                continue
            data = src.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                data = re.sub(rb"<Relationship[^>]*numbering\.xml\"/>", b"", data)
            if info.filename == "[Content_Types].xml":
                data = re.sub(rb"<Override[^>]*numbering\.xml\"[^>]*/>", b"", data)
            dst.writestr(info, data)
    raw = out.getvalue()
    assert b"numbering.xml" not in zipfile.ZipFile(io.BytesIO(raw)).read(
        "word/_rels/document.xml.rels"
    )
    prepared = docx_normalize.prepare(raw, "resume.docx", convert_bullets=True)
    document = docx.Document(io.BytesIO(prepared.raw))
    assert sum(analysis_types.is_bullet(p) for p in document.paragraphs) == 2


def test_pages_files_get_export_advice():
    with pytest.raises(docx_normalize.UploadFormatError, match="Export To"):
        docx_normalize.prepare(b"whatever", "resume.pages")


def test_other_formats_are_converted_with_libreoffice(monkeypatch, tmp_path):
    converted = _docx_bytes(_standard_resume)
    seen: list[Path] = []

    def fake_to_docx(src: Path, outdir: Path) -> Path:
        seen.append(src)
        outdir.mkdir(parents=True, exist_ok=True)
        out = outdir / f"{src.stem}.docx"
        out.write_bytes(converted)
        return out

    monkeypatch.setattr(convert, "to_docx", fake_to_docx)
    prepared = docx_normalize.prepare(b"{\\rtf1 hello}", "My Resume.rtf")
    assert seen and seen[0].suffix == ".rtf"
    assert prepared.filename == "My Resume.docx"
    assert prepared.raw == converted
    assert "converted_to_docx" in _codes(prepared)


def test_failed_conversion_is_a_readable_error(monkeypatch):
    def broken(src: Path, outdir: Path) -> Path:
        raise RuntimeError("LibreOffice binary not found")

    monkeypatch.setattr(convert, "to_docx", broken)
    with pytest.raises(docx_normalize.UploadFormatError, match="Save it as .docx"):
        docx_normalize.prepare(b"x", "resume.doc")


# --------------------------------------------------------------------------------------
# Name and contact in the page header (D5)
# --------------------------------------------------------------------------------------


def _header_identity_resume(document) -> None:
    header = document.sections[0].header
    header.paragraphs[0].text = "Ada Lovelace"
    header.add_paragraph("London • ada@example.com • (555) 010-0000")
    num_id = _add_bullet_numbering(document)
    document.add_paragraph("EDUCATION")
    document.add_paragraph("University of London | UK\t2018 - 2022")
    _make_bullet(document, "BSc Computer Science | GPA: 3.9", num_id)
    document.add_paragraph("WORK EXPERIENCES")
    document.add_paragraph("Analytical Engines | London\t2022 - Present")
    document.add_paragraph("Software Engineer")
    _make_bullet(document, "Built numerical engines in Python.", num_id)


def test_header_contact_is_kept_and_does_not_block(tmp_path):
    raw = _docx_bytes(_header_identity_resume)
    analysis = template_analyze.analyze_docx(raw=raw)
    codes = {i.code: i for i in analysis.issues}
    assert "contact_in_header" in codes and not codes["contact_in_header"].blocking
    assert "missing_contact" not in codes
    profile = analysis.suggested_profile
    assert profile is not None and analysis.ready
    assert profile.contact is None
    assert profile.name_in_header and profile.contact_in_header

    src, dst = tmp_path / "baseline.docx", tmp_path / "main_template.docx"
    src.write_bytes(raw)
    template_build.build_from_profile(src, dst, profile)
    assert template_verify.verify_tagged(dst, profile) == []
    built = docx.Document(str(dst))
    assert built.sections[0].header.paragraphs[0].text == "Ada Lovelace"
    joined = "\n".join(p.text for p in built.paragraphs)
    assert "{{ name }}" not in joined and "{{r contact }}" not in joined

    cover = tmp_path / "cover.docx"
    cover_template.build_cover_template(src, cover, profile)
    assert docx.Document(str(cover)).sections[0].header.paragraphs[0].text == "Ada Lovelace"


def test_content_import_reads_the_contact_from_the_header():
    raw = _docx_bytes(_header_identity_resume)
    document = docx.Document(io.BytesIO(raw))
    analysis = template_analyze.analyze_docx(raw=raw)
    imported = resume_import.import_from_analysis(analysis, document)
    contact = imported.resume.contact
    assert contact.name == "Ada Lovelace"
    assert contact.email == "ada@example.com"
    assert contact.phone == "(555) 010-0000"


# --------------------------------------------------------------------------------------
# Content-only import for layouts that can't be templates (D3)
# --------------------------------------------------------------------------------------


def _textbox_resume(document) -> None:
    """Two text boxes side by side: a skills sidebar on the left, experience on the right."""
    document.add_paragraph("Jamie Roe")
    document.add_paragraph("jamie@example.com")
    anchor = document.add_paragraph()

    def box(x: int, lines: list[str]) -> str:
        paras = "".join(
            f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>' for text in lines
        )
        return (
            f"<w:r {nsdecls('w', 'wp', 'a')} "
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            '<w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
            'relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{x}</wp:posOffset></wp:positionH>'
            '<wp:positionV relativeFrom="page"><wp:posOffset>900000</wp:posOffset></wp:positionV>'
            '<wp:extent cx="2000000" cy="4000000"/><wp:wrapNone/>'
            '<wp:docPr id="1" name="Text Box"/>'
            "<a:graphic><a:graphicData "
            'uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f"<wps:wsp><wps:txbx><w:txbxContent>{paras}</w:txbxContent></wps:txbx>"
            "<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>"
        )

    # Right column written first: reading order must still put the left box first.
    anchor._p.append(
        parse_xml(
            box(
                3_000_000,
                [
                    "EXPERIENCE",
                    "Gamma Labs\tJun 2025 – Aug 2025",
                    "Data Intern",
                    "• Cleaned 3 sales datasets in SQL for weekly reports",
                ],
            )
        )
    )
    anchor._p.append(parse_xml(box(400_000, ["SKILLS", "Tools: Excel, Tableau"])))


def test_textbox_layout_is_blocked_but_its_content_imports():
    raw = _docx_bytes(_textbox_resume)
    analysis = template_analyze.analyze_docx(raw=raw)
    assert any(i.code == "textboxes" and i.blocking for i in analysis.issues)

    imported = import_layout.import_content_only(docx.Document(io.BytesIO(raw)))
    sections = {s.title: s for s in imported.resume.sections}
    assert list(sections) == ["SKILLS", "EXPERIENCE"]
    job = sections["EXPERIENCE"].entries[0]
    assert (job.company, job.title, job.start) == ("Gamma Labs", "Data Intern", "2025-06")
    assert job.bullets[0].text == "Cleaned 3 sales datasets in SQL for weekly reports"
    assert imported.resume.contact.email == "jamie@example.com"
    assert "can't be used as a template" in imported.warnings[0]


def test_sidebar_table_reads_column_by_column():
    raw = _docx_bytes(_sidebar_table_resume)
    lines = [ln.text for ln in import_layout.docx_lines(docx.Document(io.BytesIO(raw)))]
    assert lines.index("JORDAN RIVERA") < lines.index("WORK EXPERIENCE")
    # The sidebar cell's bullet comes after the whole main column, not interleaved.
    assert lines.index("Improved reliability for production services.") > lines.index(
        "WORK EXPERIENCE"
    )
