"""Cover-letter settings: block toggles, letterhead choice, subject line, storage."""

from __future__ import annotations

import pytest
from docx import Document
from pydantic import ValidationError

from resume_tailor import config
from resume_tailor.content.data import CustomField
from resume_tailor.document import (
    cover_settings,
    cover_template,
    docx_text,
    render,
    template_analyze,
    template_profile,
)
from resume_tailor.document.cover_settings import CoverBlocks, CoverSettings
from resume_tailor.pipeline.coverletter_models import CoverLetter
from tests.fixtures import _docx_bytes, _full_featured_resume, _table_resume, synthetic_resume


def _letter(resume, **overrides) -> CoverLetter:
    fields = dict(
        company="Example Corp",
        company_location="Irvine, CA",
        paragraphs=["First paragraph.", "Second paragraph."],
        salutation="Dear Hiring Manager,",
        closing="Sincerely,",
        signature=resume.contact.name,
        inside_address=["Example Corp", "Irvine, CA"],
        date="August 29, 2026",
        job_title="Data Analyst",
    )
    fields.update(overrides)
    return CoverLetter(**fields)


def _render(tmp_path, build, settings: CoverSettings, resume=None, **letter):
    raw = _docx_bytes(build)
    profile = template_analyze.analyze_docx(raw=raw).suggested_profile
    src, dst = tmp_path / "baseline.docx", tmp_path / "cover.docx"
    src.write_bytes(raw)
    cover_template.build_cover_template(src, dst, profile=profile)
    resume = resume or synthetic_resume()
    out = tmp_path / "letter.docx"
    render.render_letter(
        resume,
        _letter(resume, **letter),
        template=dst,
        out=out,
        layout=template_profile.active_layout(profile),
        settings=settings,
    )
    paras = [p for p, _loc in docx_text.iter_document_paragraphs(Document(str(out)))]
    return [p.text for p in paras], {p.text: p.paragraph_format for p in paras}


def test_defaults_print_every_block_but_the_subject(tmp_path):
    texts, _fmt = _render(tmp_path, _full_featured_resume, CoverSettings())
    for line in ("August 29, 2026", "Example Corp", "Dear Hiring Manager,", "Sincerely,"):
        assert line in texts
    assert not any("Application for" in t for t in texts)


@pytest.mark.parametrize(
    "block, gone",
    [
        ("date", "August 29, 2026"),
        ("recipient", "Example Corp"),
        ("salutation", "Dear Hiring Manager,"),
        ("closing", "Sincerely,"),
    ],
)
def test_turning_a_block_off_removes_its_lines(tmp_path, block, gone):
    settings = CoverSettings(blocks=CoverBlocks(**{block: False}))
    texts, _fmt = _render(tmp_path, _full_featured_resume, settings)
    assert gone not in texts
    assert "First paragraph." in texts
    # No blank paragraph is left where the block was.
    assert texts.index("First paragraph.") - texts.index(texts[0]) < 20
    assert "" not in texts[texts.index("First paragraph.") - 2 :]


def test_subject_line_fills_title_and_company_and_takes_the_address_gap(tmp_path):
    settings = CoverSettings(
        blocks=CoverBlocks(subject=True), subject_format="Re: {job_title} at {company}"
    )
    texts, fmt = _render(tmp_path, _full_featured_resume, settings)
    assert "Re: Data Analyst at Example Corp" in texts
    assert fmt["Re: Data Analyst at Example Corp"].space_before > 0
    assert not fmt["Dear Hiring Manager,"].space_before


def test_salutation_has_no_extra_gap_when_no_address_printed(tmp_path):
    settings = CoverSettings(blocks=CoverBlocks(recipient=False))
    _texts, fmt = _render(tmp_path, _full_featured_resume, settings)
    assert not fmt["Dear Hiring Manager,"].space_before


def test_subject_with_no_title_prints_nothing(tmp_path):
    settings = CoverSettings(blocks=CoverBlocks(subject=True))
    texts, _fmt = _render(tmp_path, _full_featured_resume, settings, job_title="")
    assert not any(t.startswith("Application for") for t in texts)


def test_closing_text_overrides_the_letter(tmp_path):
    texts, _fmt = _render(tmp_path, _full_featured_resume, CoverSettings(closing_text="Best,"))
    assert "Best," in texts and "Sincerely," not in texts


def test_custom_letterhead_shows_only_the_chosen_contact_fields(tmp_path):
    resume = synthetic_resume()
    same, _ = _render(tmp_path, _full_featured_resume, CoverSettings(), resume)
    chosen, _ = _render(
        tmp_path,
        _full_featured_resume,
        CoverSettings(letterhead="custom", letterhead_fields=["email"]),
        resume,
    )
    assert any(resume.contact.phone in t for t in same)
    assert not any(resume.contact.phone in t for t in chosen)
    assert any(resume.contact.email in t for t in chosen)


def test_custom_letterhead_hides_an_unchosen_custom_line(tmp_path):
    resume = synthetic_resume().model_copy(
        update={"custom_fields": [CustomField(key="street_address", value="9 New Rd")]}
    )
    shown, _ = _render(tmp_path, _table_resume, CoverSettings(), resume)
    hidden, _ = _render(
        tmp_path, _table_resume, CoverSettings(letterhead="custom", letterhead_fields=["email"]), resume
    )
    picked, _ = _render(
        tmp_path,
        _table_resume,
        CoverSettings(letterhead="custom", letterhead_fields=["custom:street_address"]),
        resume,
    )
    cell_text = " ".join
    assert "9 New Rd" in cell_text(shown)
    assert "9 New Rd" not in cell_text(hidden)
    assert "9 New Rd" in cell_text(picked)


def test_contact_order_follows_the_template_then_canonical():
    settings = CoverSettings(letterhead="custom", letterhead_fields=["github", "phone", "email"])
    layout = {"contact_field_order": ["email", "location", "phone"]}
    assert settings.contact_fields(layout) == ["email", "phone", "github"]
    assert CoverSettings().contact_fields(layout) is None


def test_unknown_letterhead_field_is_rejected():
    with pytest.raises(ValidationError):
        CoverSettings(letterhead_fields=["fax"])


def test_load_and_save_round_trip(tmp_path, monkeypatch):
    path = tmp_path / "cover_settings.json"
    monkeypatch.setattr(config, "COVER_SETTINGS_PATH", path, raising=False)
    assert cover_settings.load() == CoverSettings()
    path.write_text("{not json", encoding="utf-8")
    assert cover_settings.load() == CoverSettings()
    saved = CoverSettings(closing_text="Warmly,", blocks=CoverBlocks(date=False))
    cover_settings.save(saved)
    assert cover_settings.load() == saved


# -- table letterhead structure ----------------------------------------------------------


def test_table_letterhead_cover_is_a_valid_word_body(tmp_path):
    """A table-layout resume keeps letterhead and sections in one table: the cover build
    used to leave cells with no paragraph and letter paragraphs after `w:sectPr`, a file
    Word refused to open ("the file appears to be corrupted")."""
    from docx.oxml.ns import qn

    raw = _docx_bytes(_table_resume)
    profile = template_analyze.analyze_docx(raw=raw).suggested_profile
    src, dst = tmp_path / "baseline.docx", tmp_path / "cover.docx"
    src.write_bytes(raw)
    cover_template.build_cover_template(src, dst, profile=profile)
    body = Document(str(dst)).element.body
    assert body[-1].tag == qn("w:sectPr")
    for cell in body.iter(qn("w:tc")):
        assert cell.find(qn("w:p")) is not None
    # Only the letterhead rows survive; the resume sections are gone.
    texts = [p.text for p, _loc in docx_text.iter_document_paragraphs(Document(str(dst)))]
    assert not any("EDUCATION" in t.upper() for t in texts)


def test_table_letterhead_keeps_side_margins_and_aligns_the_body(tmp_path):
    raw = _docx_bytes(_table_resume)
    profile = template_analyze.analyze_docx(raw=raw).suggested_profile
    src, dst = tmp_path / "baseline.docx", tmp_path / "cover.docx"
    src.write_bytes(raw)
    cover_template.build_cover_template(src, dst, profile=profile)
    original, cover = Document(str(src)).sections[0], Document(str(dst)).sections[0]
    assert cover.left_margin == original.left_margin
    date = next(
        p for p, _loc in docx_text.iter_document_paragraphs(Document(str(dst)))
        if p.text == "{{ date }}"
    )
    assert date.paragraph_format.left_indent > 0  # at least the cell padding
