"""Education prototype tagging: a plain degree line must not become a bullet."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn

from resume_tailor.document import (
    template_analyze,
    template_build,
    template_bullets,
    template_tags,
)
from tests.fixtures import _docx_bytes, _table_resume


def _degree_paragraph(path: Path):
    doc = docx.Document(str(path))
    for table in doc.tables:
        for row in table.rows:
            for tc in row._tr.findall(qn("w:tc")):
                for p in tc.findall(qn("w:p")):
                    text = "".join(t.text or "" for t in p.iter(qn("w:t")))
                    if text == template_tags.EDUCATION_DEGREE_TAG:
                        return p
    raise AssertionError("degree tag not found")


def _mark_size(p) -> str | None:
    sz = p.find(f"{qn('w:pPr')}/{qn('w:rPr')}/{qn('w:sz')}")
    return sz.get(qn("w:val")) if sz is not None else None


def test_plain_table_degree_line_gets_no_numbering(tmp_path: Path, monkeypatch):
    """The fixture's degree line ("BS Computer Science") is plain italic text in the
    header cell. Tagging used to run `retarget_bullet` on it unconditionally, which
    attached the canonical bullet list (marker + hanging indent) and shrank its
    paragraph mark as if it were a bullet.

    The fixture's list has no Noto Sans Symbols marker, so discovery is pinned to it —
    a real Google Docs export (the reported case) always has one, which is what makes
    `retarget_bullet` attach a list to a paragraph that had none."""
    monkeypatch.setattr(template_bullets, "discover_noto_num_id", lambda doc: "20")
    raw = _docx_bytes(_table_resume)
    analysis = template_analyze.analyze_docx(raw=raw)
    profile = analysis.suggested_profile
    assert profile is not None and profile.layout == "table"
    src, dst = tmp_path / "baseline.docx", tmp_path / "main_template.docx"
    src.write_bytes(raw)

    before = docx.Document(str(src))
    original = next(
        p._p for t in before.tables for c in t._cells for p in c.paragraphs
        if p.text == "BS Computer Science"
    )
    template_build.build_from_profile(src, dst, profile)

    degree = _degree_paragraph(dst)
    assert degree.find(f".//{qn('w:numPr')}") is None
    assert _mark_size(degree) == _mark_size(original)
