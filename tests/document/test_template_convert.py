"""`template_convert.to_generic`: a saved fixed-mode profile gains the generic-only fields
from its baseline, keeping every confirmed mapping."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.shared import Pt

from resume_tailor.document import template_analyze, template_build, template_convert
from tests.fixtures import _docx_bytes, _standard_resume, as_fixed


def _fixed(tmp_path: Path, build=_standard_resume):
    raw = _docx_bytes(build)
    baseline = tmp_path / "original_export.docx"
    baseline.write_bytes(raw)
    suggested = template_analyze.analyze_docx(raw=raw).suggested_profile
    assert suggested is not None
    return baseline, suggested, as_fixed(suggested)


def test_to_generic_matches_what_a_fresh_import_suggests(tmp_path):
    baseline, suggested, fixed = _fixed(tmp_path)
    converted = template_convert.to_generic(fixed, baseline)

    assert converted.section_mode == "generic"
    assert converted.sections == suggested.sections
    assert converted.heading_prototype == suggested.heading_prototype
    assert converted.spacing == suggested.spacing
    # Confirmed mappings are carried over untouched.
    for attr in ("experience", "education", "projects", "skills", "contact", "enabled"):
        assert getattr(converted, attr) == getattr(fixed, attr)


def test_converted_profile_builds_a_generic_template(tmp_path):
    baseline, _, fixed = _fixed(tmp_path)
    converted = template_convert.to_generic(fixed, baseline)
    dst = tmp_path / "main_template.docx"
    template_build.build_from_profile(baseline, dst, converted)
    joined = "\n".join(p.text for p in docx.Document(str(dst)).paragraphs)
    assert "{%p for section in sections %}" in joined
    assert "{{ section.title }}" in joined


def test_generic_profile_passes_through(tmp_path):
    baseline, suggested, _ = _fixed(tmp_path)
    assert template_convert.to_generic(suggested, baseline) is suggested


def test_no_enabled_section_is_refused(tmp_path):
    baseline, _, fixed = _fixed(tmp_path)
    nothing = fixed.model_copy(
        update={
            "enabled": fixed.enabled.model_copy(
                update={"experience": False, "education": False, "projects": False, "skills": False}
            )
        }
    )
    with pytest.raises(template_convert.ConversionError):
        template_convert.to_generic(nothing, baseline)


def test_heading_differences_become_warnings(tmp_path):
    def build(document):
        _standard_resume(document)
        for paragraph in document.paragraphs:
            if paragraph.text == "SKILLS":
                for run in paragraph.runs:
                    run.font.size = Pt(16)

    baseline, _, fixed = _fixed(tmp_path, build)
    converted = template_convert.to_generic(fixed.model_copy(update={"warnings": []}), baseline)
    flagged = [w for w in converted.warnings if w.startswith("Heading ")]
    assert len(flagged) == 1
    assert "'SKILLS'" in flagged[0] and "text formatting" in flagged[0]
