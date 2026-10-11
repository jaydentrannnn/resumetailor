"""The Edit-layout-in-Word skeleton: make it, edit it, read it back."""

from __future__ import annotations

import copy
import io
import zipfile
from pathlib import Path

import docx
import pytest
from docx.text.paragraph import Paragraph

from resume_tailor import config
from resume_tailor.content.data import CustomField
from resume_tailor.document import (
    date_style,
    docx_text,
    render,
    skeleton_parse,
    template_analyze,
    template_build,
    template_profile,
    template_skeleton,
)
from tests.fixtures import _docx_bytes, _full_featured_resume, _table_resume, synthetic_resume


@pytest.fixture
def built(tmp_path, monkeypatch):
    """(baseline path, tagged template, profile) for a fixture builder."""

    def make(builder):
        raw = _docx_bytes(builder)
        profile = template_analyze.analyze_docx(raw=raw).suggested_profile
        src, tpl = tmp_path / "baseline.docx", tmp_path / "tagged.docx"
        src.write_bytes(raw)
        template_build.build_from_profile(src, tpl, profile)
        monkeypatch.setattr(config, "BASELINE_TEMPLATE_PATH", src)
        return src, tpl, profile

    return make


def _paras(raw: bytes):
    doc = docx.Document(io.BytesIO(raw))
    return doc, [p for p, _loc in docx_text.iter_document_paragraphs(doc)]


def _save(doc) -> bytes:
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _set(paragraph: Paragraph, text: str) -> None:
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    paragraph.runs[0].text = text


def _rebuild(tmp_path: Path, parsed, monkeypatch, resume=None) -> list[str]:
    """Install the read-back as a template and render `resume` through it."""
    assert parsed.result.ready, parsed.result.issues
    profile = parsed.result.suggested_profile
    src, tpl = tmp_path / "realised.docx", tmp_path / "tagged2.docx"
    src.write_bytes(parsed.raw)
    monkeypatch.setattr(config, "BASELINE_TEMPLATE_PATH", src)
    template_build.build_from_profile(src, tpl, profile)
    out = tmp_path / "out.docx"
    render.render(
        resume or synthetic_resume(), template=tpl, out=out,
        layout=template_profile.active_layout(profile),
    )  # fmt: skip
    return [p.text for p, _loc in docx_text.iter_document_paragraphs(docx.Document(str(out)))]


def _render(tpl: Path, profile, tmp_path: Path) -> list[str]:
    out = tmp_path / "before.docx"
    render.render(synthetic_resume(), template=tpl, out=out,
                  layout=template_profile.active_layout(profile))  # fmt: skip
    return [p.text for p, _loc in docx_text.iter_document_paragraphs(docx.Document(str(out)))]


@pytest.mark.parametrize("builder", [_full_featured_resume, _table_resume])
def test_an_untouched_layout_file_reads_back_to_the_same_resume(built, builder, tmp_path, monkeypatch):
    src, tpl, profile = built(builder)
    before = _render(tpl, profile, tmp_path)
    style = date_style.detect(src) or date_style.DEFAULT
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    parsed = skeleton_parse.parse_skeleton(skeleton.raw, skeleton.meta(), synthetic_resume())
    assert _rebuild(tmp_path, parsed, monkeypatch) == before
    # The realised baseline keeps the template's date style, which is read from its text.
    (tmp_path / "check.docx").write_bytes(parsed.raw)
    assert (date_style.detect(tmp_path / "check.docx") or date_style.DEFAULT) == style


def test_skeleton_shows_tokens_bookmarks_and_notes_but_no_tags(built):
    _src, tpl, profile = built(_full_featured_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    _doc, paras = _paras(skeleton.raw)
    texts = [p.text for p in paras]
    assert texts[0] == "«Name»"
    assert "«Company» | «Job location»\t«Job dates»" in texts
    xml = zipfile.ZipFile(io.BytesIO(skeleton.raw)).read("word/document.xml").decode()
    assert "{{" not in xml and "{%" not in xml and "QQ" not in xml
    assert 'w:name="rt_sec_0"' in xml
    assert "word/comments.xml" in zipfile.ZipFile(io.BytesIO(skeleton.raw)).namelist()
    assert [s.kind for s in skeleton.sections] == ["education", "experience", "project", "skills"]


def test_split_bolded_and_moved_tokens_are_read_exactly(built, tmp_path, monkeypatch):
    _src, tpl, profile = built(_full_featured_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    doc, paras = _paras(skeleton.raw)
    header = next(p for p in paras if p.text.startswith("«Company»"))
    title = paras[paras.index(header) + 1]
    # Word splits a token across runs when part of it is restyled.
    _set(header, "«Comp")
    header.add_run("any»").italic = True
    header.add_run(" | «Job location»")
    # Dates move from the company line to the title line.
    title.add_run("\t«Job dates»")
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta(), synthetic_resume())
    exp = parsed.result.suggested_profile.experience
    assert exp.header.date_alignment == "separate_paragraph"
    assert exp.header.date_paragraph_id == exp.title_paragraph_id
    texts = _rebuild(tmp_path, parsed, monkeypatch)
    company = synthetic_resume().experience[0]
    assert f"{company.company} | {company.location}" in texts
    assert any(t.startswith(f"{company.title}\t") for t in texts)


def test_a_deleted_optional_field_is_dropped_and_a_required_one_blocks(built):
    _src, tpl, profile = built(_full_featured_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    doc, paras = _paras(skeleton.raw)
    header = next(p for p in paras if p.text.startswith("«Company»"))
    _set(header, "«Company»\t«Job dates»")
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta())
    assert parsed.result.ready
    assert "location" not in parsed.result.suggested_profile.experience.header.fields

    for p in paras:
        if p.text == "«Job bullet»":
            _set(p, "")
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta())
    assert not parsed.result.ready
    assert "layout_field_missing" in {i.code for i in parsed.result.issues}


def test_a_typed_custom_line_becomes_an_editable_line(built, tmp_path, monkeypatch):
    _src, tpl, profile = built(_full_featured_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    doc, paras = _paras(skeleton.raw)
    line = copy.deepcopy(paras[1]._p)
    paras[1]._p.addnext(line)
    _set(Paragraph(line, paras[1]._parent), "<<custom: Portfolio>>")
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta())
    [slot] = parsed.result.suggested_profile.custom_fields
    assert (slot.key, slot.label) == ("portfolio", "Portfolio")
    resume = synthetic_resume().model_copy(
        update={"custom_fields": [CustomField(key="portfolio", value="ada.dev")]}
    )
    assert "ada.dev" in _rebuild(tmp_path, parsed, monkeypatch, resume)


def test_table_custom_line_keeps_its_key_and_text(built):
    _src, tpl, profile = built(_table_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    _doc, paras = _paras(skeleton.raw)
    assert "«custom: Street address»" in [p.text for p in paras]
    parsed = skeleton_parse.parse_skeleton(skeleton.raw, skeleton.meta())
    [slot] = parsed.result.suggested_profile.custom_fields
    assert (slot.key, slot.text) == ("street_address", "123 Example St.")
    assert parsed.result.suggested_profile.layout == "table"


def test_retyped_heading_is_found_by_text_and_a_deleted_one_is_reported(built):
    _src, tpl, profile = built(_full_featured_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    doc, paras = _paras(skeleton.raw)
    heading = next(p for p in paras if p.text == "PROJECTS")
    for mark in heading._p.findall(".//{*}bookmarkStart") + heading._p.findall(".//{*}bookmarkEnd"):
        mark.getparent().remove(mark)
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta())
    assert parsed.result.suggested_profile.projects is not None

    _set(heading, "Side work")
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta())
    assert "layout_section_missing" in {i.code for i in parsed.result.issues}


def test_unknown_token_is_reported_and_comments_are_stripped(built):
    _src, tpl, profile = built(_full_featured_resume)
    skeleton = template_skeleton.make_skeleton(tpl, profile, synthetic_resume())
    doc, paras = _paras(skeleton.raw)
    paras[1].add_run(" «Fax»")
    parsed = skeleton_parse.parse_skeleton(_save(doc), skeleton.meta())
    assert "layout_unknown_token" in {i.code for i in parsed.result.issues}
    names = zipfile.ZipFile(io.BytesIO(parsed.raw)).namelist()
    xml = zipfile.ZipFile(io.BytesIO(parsed.raw)).read("word/document.xml").decode()
    assert "commentReference" not in xml and "commentRangeStart" not in xml
    rels = zipfile.ZipFile(io.BytesIO(parsed.raw)).read("word/_rels/document.xml.rels").decode()
    assert "comments" not in rels
    assert names  # still a valid package
