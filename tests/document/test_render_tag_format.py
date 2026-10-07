"""A `{{r }}` tag's run formatting survives the RichText that replaces it."""

import io
import re
import zipfile

from docx import Document
from docx.shared import Pt
from docxtpl import DocxTemplate

from resume_tailor.content.data import MasterResume
from resume_tailor.document import render


def _template(font: str, size: int) -> DocxTemplate:
    doc = Document()
    run = doc.add_paragraph().add_run("{{r contact }}")
    run.font.name, run.font.size, run.bold = font, Pt(size), True
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return DocxTemplate(buffer)


def test_tag_formats_read_the_run_the_tag_sits_in():
    assert render._tag_formats(_template("Times New Roman", 12))["contact"] == {
        "font": "Times New Roman", "size": 24, "bold": True,
    }


def test_contact_richtext_restates_the_run_font_and_size():
    resume = MasterResume.model_validate({
        "contact": {"name": "Ada", "email": "ada@example.com", "phone": "555", "location": "Irvine, CA"},
        "sections": [],
    })
    tpl = _template("Times New Roman", 12)
    rich = render._contact_richtext(resume, tpl)
    tpl.render({"contact": rich})
    out = io.BytesIO()
    tpl.save(out)
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    runs = re.findall(r"<w:r>.*?</w:r>", xml, re.S)
    texted = [run for run in runs if "ada@example.com" in run]
    assert texted
    assert 'w:ascii="Times New Roman"' in texted[0]
    assert '<w:sz w:val="24"/>' in texted[0]
