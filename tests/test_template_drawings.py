"""Decorative drawings (B15): a photo, an icon or a rule never blocks a template, and
survives the build and the render untouched. Text boxes holding text still block.
"""

from __future__ import annotations

import io
import struct
import zipfile
import zlib

from resume_tailor import render, template_analyze, template_build
from tests.fixtures import _docx_bytes, _full_featured_resume, synthetic_resume


def _png() -> bytes:
    """A valid 1x1 PNG, built in memory."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    pixels = zlib.compress(b"\x00\x00\x00\x00\x00")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(
        b"IEND", b""
    )


def _with_drawings(document) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    _full_featured_resume(document)
    # An icon beside the name, and a standalone picture paragraph (a divider graphic).
    document.paragraphs[0].add_run().add_picture(io.BytesIO(_png()), width=Pt(10))
    # A phone glyph *before* the contact text: the tag must land in a text run.
    contact = document.paragraphs[1]
    glyph = contact.add_run()
    glyph.add_picture(io.BytesIO(_png()), width=Pt(8))
    first_run = contact._p.find(qn("w:r"))
    first_run.addprevious(glyph._r)
    divider = document.paragraphs[1].insert_paragraph_before()
    divider.add_run().add_picture(io.BytesIO(_png()), width=Pt(400), height=Pt(1))


def _drawings(docx_path) -> int:
    with zipfile.ZipFile(docx_path) as archive:
        return archive.read("word/document.xml").count(b"<w:drawing>")


def test_drawings_are_a_non_blocking_note(tmp_path):
    result = template_analyze.analyze_docx(raw=_docx_bytes(_with_drawings))
    codes = {issue.code: issue.blocking for issue in result.issues}
    assert codes.get("decorative_drawing") is False
    assert "textboxes" not in codes
    assert result.ready, result.issues


def test_drawings_survive_build_and_render(tmp_path):
    src = tmp_path / "original_export.docx"
    src.write_bytes(_docx_bytes(_with_drawings))
    result = template_analyze.analyze_docx(raw=src.read_bytes())
    built = tmp_path / "main_template.docx"
    template_build.build_from_profile(src, built, result.suggested_profile)
    assert _drawings(built) == _drawings(src) == 3
    out = tmp_path / "out.docx"
    render.render(synthetic_resume(), template=built, out=out)
    assert _drawings(out) == 3
    with zipfile.ZipFile(out) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    resume = synthetic_resume()
    assert resume.contact.name in xml and resume.contact.email in xml


def _with_text_box(document) -> None:
    from docx.oxml import parse_xml

    _full_featured_resume(document)
    ns = (
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"'
    )
    run = parse_xml(
        f"<w:r {ns}><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>"
        "Skills sidebar</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>"
    )
    document.paragraphs[1]._p.append(run)


def test_text_in_a_text_box_still_blocks():
    result = template_analyze.analyze_docx(raw=_docx_bytes(_with_text_box))
    blocking = {issue.code for issue in result.issues if issue.blocking}
    assert "textboxes" in blocking
