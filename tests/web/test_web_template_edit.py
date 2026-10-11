"""`/api/template/edit-session`: Edit layout in Word, from writing the file to install."""

from __future__ import annotations

import io
import json
import os
from pathlib import Path

import docx
import pytest

from resume_tailor import config
from resume_tailor.document import docx_text, template_analyze, template_build, template_profile
from resume_tailor.web import state as web_state
from resume_tailor.web import template_edit, template_install
from tests.fixtures import _docx_bytes, _full_featured_resume
from tests.web.helpers import _DOCX_MIME, _point_templates_at


@pytest.fixture
def desktop(client, tmp_path, monkeypatch):
    """A desktop app with an installed template; returns (client, opened paths)."""
    c, _ = client
    templates = _point_templates_at(tmp_path, monkeypatch)
    raw = _docx_bytes(_full_featured_resume)
    profile = template_analyze.analyze_docx(raw=raw).suggested_profile
    (templates / "original_export.docx").write_bytes(raw)
    template_profile.save_profile(profile, templates / "template_profile.json")
    template_build.build_from_profile(
        templates / "original_export.docx", templates / "main_template.docx", profile
    )
    monkeypatch.setattr(web_state, "desktop", True)
    opened: list[Path] = []
    monkeypatch.setattr(template_edit, "_open_file", opened.append)
    return c, opened


def _save_in_word(edit) -> None:
    """Underline «Company» and save, moving the file's mtime the way a real save does."""
    path = template_edit.layout_path()
    doc = docx.Document(str(path))
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if "«Company»" in run.text:
                run.underline = True
    buffer = io.BytesIO()
    doc.save(buffer)
    path.write_bytes(buffer.getvalue())
    stat = path.stat()
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 2_000_000_000))
    edit(path)


def test_session_is_desktop_only(client, monkeypatch):
    c, _ = client
    monkeypatch.setattr(web_state, "desktop", False)
    assert c.get("/api/config").json()["desktop"] is False
    assert c.get("/api/template/edit-session").json() == {
        "active": False, "changed": False, "saved_at": None, "editor_open": False,
        "editor_closed": False, "started_at": None, "path": None,
    }  # fmt: skip
    assert c.post("/api/template/edit-session").status_code == 404


def test_start_writes_and_opens_the_layout_file_once(desktop):
    c, opened = desktop
    res = c.post("/api/template/edit-session")
    assert res.status_code == 200, res.json()
    assert res.json()["active"] and not res.json()["changed"]
    assert opened == [template_edit.layout_path()]
    assert docx.Document(str(opened[0])).paragraphs[0].text == "«Name»"
    assert c.post("/api/template/edit-session").status_code == 409

    assert c.post("/api/template/edit-session/open").status_code == 200
    assert len(opened) == 2


def test_status_notices_a_save_and_the_editor_lock_file(desktop):
    c, _opened = desktop
    c.post("/api/template/edit-session")
    lock = template_edit.layout_path().parent / "~$yout.docx"
    lock.write_bytes(b"owner")
    assert c.get("/api/template/edit-session").json()["editor_open"] is True
    _save_in_word(lambda _p: lock.unlink())
    body = c.get("/api/template/edit-session").json()
    assert body["changed"] is True and body["saved_at"]
    assert body["editor_open"] is False and body["editor_closed"] is True


def test_status_tells_closed_without_saving_from_not_opened_yet(desktop):
    c, _opened = desktop
    c.post("/api/template/edit-session")
    assert c.get("/api/template/edit-session").json()["editor_closed"] is False
    lock = template_edit.layout_path().parent / "~$yout.docx"
    lock.write_bytes(b"owner")
    c.get("/api/template/edit-session")
    lock.unlink()
    body = c.get("/api/template/edit-session").json()
    assert body["editor_closed"] is True and body["changed"] is False
    c.post("/api/template/edit-session/open")  # opening again starts the wait over
    assert c.get("/api/template/edit-session").json()["editor_closed"] is False


def test_saved_layout_reads_back_and_installs(desktop, monkeypatch):
    c, _opened = desktop
    c.post("/api/template/edit-session")
    _save_in_word(lambda _p: None)
    analysis = c.post("/api/template/edit-session/import")
    assert analysis.status_code == 200, analysis.json()
    body = analysis.json()
    assert body["ready"], body["issues"]
    sha = body["source_sha256"]
    realised = c.get(f"/api/template/drafts/{sha}")
    assert realised.status_code == 200
    assert b"\xc2\xab" not in realised.content  # no « left in the package bytes' text

    def fake_build(*, source=None, output=None, profile_path=None, legacy=False):
        out = Path(output) if output else config.DEFAULT_TEMPLATE_PATH
        out.parent.mkdir(parents=True, exist_ok=True)
        loaded = template_profile.load_profile(Path(profile_path))
        template_build.build_from_profile(Path(source), out, loaded)
        return 0, "stub build ok"

    monkeypatch.setattr(template_install, "_run_build", fake_build)
    install = c.post(
        "/api/template",
        data={"profile": json.dumps(body["suggested_profile"])},
        files={"file": ("layout.docx", realised.content, _DOCX_MIME)},
    )
    assert install.status_code == 200, install.json()
    # The installed baseline is the read-back, with the underlined company.
    header = body["suggested_profile"]["experience"]["header"]["header_paragraph_id"]
    paragraphs = [p for p, _loc in docx_text.iter_document_paragraphs(
        docx.Document(str(config.BASELINE_TEMPLATE_PATH)))]  # fmt: skip
    assert paragraphs[header].runs[0].underline

    assert c.delete("/api/template/edit-session").status_code == 204
    assert c.get("/api/template/edit-session").json()["active"] is False


def test_import_without_a_session_or_with_a_broken_file_is_refused(desktop):
    c, _opened = desktop
    assert c.post("/api/template/edit-session/import").status_code == 409
    c.post("/api/template/edit-session")
    template_edit.layout_path().write_bytes(b"half-written")
    res = c.post("/api/template/edit-session/import")
    assert res.status_code == 409
    assert "still being saved" in res.json()["detail"]


def test_quick_adjustments_read_the_layout_and_return_an_installable_draft(desktop):
    c, opened = desktop
    form = c.get("/api/template/adjustments").json()
    current = form["current"]
    assert form["table_layout"] is False
    assert current["margins"]["left"] > 0 and current["font_size"] > 0
    company = next(f for f in current["fields"] if f["label"] == "Company")
    assert company["bold"] is True

    res = c.post(
        "/api/template/adjustments",
        json={
            "margins": {**current["margins"], "left": 0.5, "right": 0.5},
            "font_family": "Georgia",
            "font_size": current["font_size"] + 1,
            "fields": [{"label": "Company", "bold": False, "italic": True}],
        },
    )
    assert res.status_code == 200, res.json()
    body = res.json()
    assert body["ready"], body["issues"]
    draft = docx.Document(io.BytesIO(c.get(f"/api/template/drafts/{body['source_sha256']}").content))
    assert round(draft.sections[0].left_margin.inches, 2) == 0.5
    header = body["suggested_profile"]["experience"]["header"]["header_paragraph_id"]
    paragraphs = [p for p, _loc in docx_text.iter_document_paragraphs(draft)]
    run = paragraphs[header].runs[0]
    assert run.italic and not run.bold and run.font.name == "Georgia"
    assert opened == []  # no word processor involved
