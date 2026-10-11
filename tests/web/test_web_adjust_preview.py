"""`/api/template/adjustments/preview`: the live Quick adjustments preview."""

from __future__ import annotations

import docx
import pytest

from resume_tailor import config
from resume_tailor.document import convert_warm, template_analyze, template_build, template_profile
from resume_tailor.document import thumbnails
from resume_tailor.web import adjust_preview, template_edit
from tests.fixtures import _docx_bytes, _full_featured_resume
from tests.web.helpers import _point_templates_at


@pytest.fixture
def previews(client, tmp_path, monkeypatch):
    """An installed template, a cache dir, and a fake PDF engine that records the
    rendered documents (one page each)."""
    c, _ = client
    templates = _point_templates_at(tmp_path, monkeypatch)
    raw = _docx_bytes(_full_featured_resume)
    profile = template_analyze.analyze_docx(raw=raw).suggested_profile
    (templates / "original_export.docx").write_bytes(raw)
    template_profile.save_profile(profile, templates / "template_profile.json")
    template_build.build_from_profile(
        templates / "original_export.docx", templates / "main_template.docx", profile
    )
    monkeypatch.setattr(config, "CACHE_DIR", tmp_path / "cache")
    rendered: list = []

    def to_pdf(docx_path, pdf_path):
        rendered.append(docx.Document(str(docx_path)))
        pdf_path.write_bytes(b"%PDF-fake")
        return pdf_path

    def pages(_pdf, directory, *, width):
        (directory / "1.png").write_bytes(b"\x89PNG fake")
        return 1

    monkeypatch.setattr(convert_warm, "to_pdf", to_pdf)
    monkeypatch.setattr(thumbnails, "pdf_pages_png", pages)
    return c, rendered


def test_preview_renders_the_resume_through_the_adjusted_template(previews):
    c, rendered = previews
    body = {"separators": {"tech_items": " / "}, "font_family": "Georgia"}
    res = c.post("/api/template/adjustments/preview", json=body)
    assert res.status_code == 200, res.json()
    key, pages = res.json()["key"], res.json()["pages"]
    assert pages == 1
    text = "\n".join(p.text for p in rendered[0].paragraphs)
    assert "Python / FastAPI" in text
    assert all(r.font.name == "Georgia" for p in rendered[0].paragraphs for r in p.runs if r.text)
    page = c.get(f"/api/template/adjustments/preview/{key}/1.png")
    assert page.status_code == 200 and page.headers["content-type"] == "image/png"
    # The same adjustments come from the cache: no second conversion.
    assert c.post("/api/template/adjustments/preview", json=body).json()["key"] == key
    assert len(rendered) == 1
    # Nothing was installed.
    assert template_profile.load_profile().style.separators.tech_items == ", "


def test_a_newer_request_supersedes_one_that_has_not_converted_yet(previews, monkeypatch):
    c, rendered = previews
    real = template_edit.adjusted

    def adjusted_then_newer_request(adj):
        out = real(adj)
        adjust_preview._take_ticket()  # another preview arrived meanwhile
        return out

    monkeypatch.setattr(template_edit, "adjusted", adjusted_then_newer_request)
    res = c.post("/api/template/adjustments/preview", json={"font_size": 12})
    assert res.status_code == 409 and res.json()["detail"] == "superseded"
    assert rendered == []


def test_unknown_pages_and_bad_keys_are_not_found(previews):
    c, _rendered = previews
    assert c.get("/api/template/adjustments/preview/../1.png").status_code == 404
    assert c.get(f"/api/template/adjustments/preview/{'a' * 24}/1.png").status_code == 404
    assert c.get(f"/api/template/adjustments/preview/{'a' * 24}/0.png").status_code == 404


def test_no_template_is_a_clear_error(client, tmp_path, monkeypatch):
    c, _ = client
    _point_templates_at(tmp_path, monkeypatch)
    res = c.post("/api/template/adjustments/preview", json={})
    assert res.status_code == 422
    assert "template" in res.json()["detail"].lower()
