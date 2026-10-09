"""Shared helpers for the `tests/web/test_web_*.py` files."""

from __future__ import annotations

import json
import time
from pathlib import Path

import pytest

from resume_tailor import config
from resume_tailor.pipeline import jd
from tests.fixtures import synthetic_resume


def _stub_extract_consensus(text, *, known_tags=None, runs=1, use_cache=True, on_event=None):
    """Route `extract_consensus` straight to `jd.extract` for web tests.

    `jobs.py` now calls `jd.extract_consensus`, which at `runs > 1` (the default) would
    call `jd.extract` multiple times and write a real consensus cache file. These tests
    stub `jd.extract` directly and care about the job/queue contract, not the voting
    algorithm (that's `test_jd.py`'s job) — a live attribute lookup on `jd.extract`
    means each test's own stub is still honoured.
    """
    return jd.extract(text, known_tags=known_tags, use_cache=use_cache, on_event=on_event)


def _disk_only_run(job_id: str) -> Path:
    """Write a finished run's `run.json` straight to disk; returns its folder."""
    jobs_dir = config.OUTPUT_DIR / "jobs" / job_id
    jobs_dir.mkdir(parents=True, exist_ok=True)
    (jobs_dir / "run.json").write_text(
        json.dumps(
            {
                "job_id": job_id,
                "workspace_id": config.active_workspace_id(),
                "created_at": "2026-01-02T00:00:00+00:00",
                "finished_at": "2026-01-02T00:01:00+00:00",
                "status": "succeeded",
                "title": "Disk Only",
                "error": None,
                "report": {"title": "Disk Only", "seniority": "intern",
                           "coverage_matched": 0, "coverage_total": 0,
                           "missing_must_haves": [], "unmatched_canonicals": [],
                           "gaps": [], "model": "stub", "semantic_used": False,
                           "bullets_selected": 0, "bullets_total": 0,
                           "experience": [], "projects": [], "dropped": [],
                           "pages": 1, "pages_are_estimated": True, "iterations": 1,
                           "widows_repaired": 0, "widows_remaining": 0,
                           "verbs_diversified": 0, "verb_collisions_remaining": 0,
                           "warnings": [], "out_path": str(jobs_dir / "tailored.docx"),
                           "pdf_backend": "soffice", "calibration_source": "fallback"},
            }
        ),
        encoding="utf-8",
    )
    return jobs_dir


def _point_settings_at(tmp_path: Path, monkeypatch) -> Path:
    """Redirect settings.json under tmp_path so a test never touches the real repo."""
    path = tmp_path / "settings.json"
    monkeypatch.setattr(config, "SETTINGS_PATH", path)
    return path


def _stub_no_network_extract(*args, **kwargs):
    """Fail the run at its first pipeline call, after routing has been resolved."""
    raise RuntimeError("stub — no network in tests")


def _drain(c, job_id: str) -> dict:
    """Block until the queue's worker thread has finished `job_id`, and return its status.

    The worker is a background thread, so anything asserting on what a *run* did (rather
    than on the queued `Job` object) has to wait for it or it is a race that passes on a
    fast machine.
    """

    deadline = time.time() + 10
    while time.time() < deadline:
        status = c.get(f"/api/jobs/{job_id}").json()
        if status["status"] in ("succeeded", "failed"):
            return status
        time.sleep(0.05)
    pytest.fail("job did not finish in time")


class _FakeSDKError(Exception):
    """Stands in for `anthropic.BadRequestError` (an `Exception`, not a `RuntimeError`)
    without pulling the real SDK's exception hierarchy into this hermetic suite."""


def _minimal_docx_bytes(paragraph: str | None = None) -> bytes:
    """Build a tiny valid .docx in memory for upload tests (no Word required).

    Optional `paragraph` text makes two fixtures differ byte-for-byte.
    """
    import io

    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    if paragraph is not None:
        doc.add_paragraph(paragraph)
    doc.save(buf)
    return buf.getvalue()


def _resume_docx_bytes() -> bytes:
    """A single-column resume the analyzer accepts (`ready=True`).

    `_minimal_docx_bytes` is deliberately contentless and yields no suggested profile,
    so it cannot drive the wizard/profile install path.
    """
    import io

    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()

    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "9")
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    lvl.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "●")
    lvl.append(lvl_text)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "9")
    abs_el = OxmlElement("w:abstractNumId")
    abs_el.set(qn("w:val"), "9")
    num.append(abs_el)
    numbering.append(num)

    def bullet(text: str):
        paragraph = document.add_paragraph(text)
        pPr = paragraph._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), "9")
        numPr.append(ilvl)
        numPr.append(nid)
        pPr.append(numPr)

    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("London • ada@example.com • LinkedIn")
    document.add_paragraph("EDUCATION")
    document.add_paragraph("University of London | UK\t2018 - 2022")
    bullet("BSc Computer Science | GPA: 3.9")
    bullet("Relevant Coursework: Algorithms, Databases")
    document.add_paragraph("WORK EXPERIENCES")
    document.add_paragraph("Analytical Engines | London\t2022 - Present")
    document.add_paragraph("Software Engineer")
    bullet("Built numerical engines in Python.")
    document.add_paragraph("PROJECTS")
    document.add_paragraph("Note Engine | Python, FastAPI\t2024")
    bullet("Indexed research notes with embeddings.")
    document.add_paragraph("SKILLS")
    skills = document.add_paragraph()
    skills.add_run("Languages:").bold = True
    skills.add_run(" Python, SQL")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _resume_upload_with_profile() -> tuple[bytes, str]:
    """A real, analyzable resume upload plus its suggested profile as a JSON string —
    `POST /api/template` requires a profile now that there is no legacy hard-coded-
    heading fallback (Phase 7 retired it), so any test exercising the install plumbing
    itself (backup/restore, library recording, busy-queue rejection, …) needs a real
    uploadable file and a profile that actually matches it, not `_minimal_docx_bytes`."""
    from resume_tailor.document import template_analyze

    raw = _resume_docx_bytes()
    result = template_analyze.analyze_docx(raw=raw)
    assert result.suggested_profile is not None, result.issues
    return raw, result.suggested_profile.model_dump_json()


def _point_templates_at(tmp_path: Path, monkeypatch) -> Path:
    """Redirect baseline/tagged/library paths under tmp_path and return the templates dir."""
    templates = tmp_path / "templates"
    templates.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(config, "TEMPLATES_DIR", templates)
    monkeypatch.setattr(config, "BASELINE_TEMPLATE_PATH", templates / "original_export.docx")
    monkeypatch.setattr(config, "DEFAULT_TEMPLATE_PATH", templates / "main_template.docx")
    monkeypatch.setattr(config, "TEMPLATE_PROFILE_PATH", templates / "template_profile.json")
    monkeypatch.setattr(config, "TEMPLATE_LIBRARY_DIR", templates / "library")
    return templates


_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _write_test_resume(monkeypatch, tmp_path, *, bullet_text: str, bullet_tags: list[str]) -> None:
    path = tmp_path / "test_master_resume.json"
    path.write_text(
        json.dumps(
            {
                "contact": {"name": "X", "email": "x@y.z"},
                "experience": [
                    {
                        "company": "Acme",
                        "title": "Engineer",
                        "start": "2020",
                        "end": "2021",
                        "bullets": [{"id": "b1", "text": bullet_text, "tags": bullet_tags}],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(config, "MASTER_RESUME_PATH", path)


class _FakeProposeResponse:
    def __init__(self, parsed):
        self.parsed_output = parsed
        self.stop_reason = "end_turn"


class _FakeProposeMessages:
    def __init__(self, parsed, calls):
        self._parsed = parsed
        self._calls = calls

    def parse(self, **kwargs):
        self._calls.append(kwargs)
        return _FakeProposeResponse(self._parsed)


class _FakeProposeClient:
    def __init__(self, parsed, calls):
        self.messages = _FakeProposeMessages(parsed, calls)


def _point_workspaces_at(tmp_path: Path, monkeypatch) -> dict[str, Path]:
    """Redirect every workspace storage root under tmp_path so tests never touch the
    real repo's data/templates/output trees, and reset the active-workspace pointer."""
    data_root = tmp_path / "ws_data"
    templates_root = tmp_path / "ws_templates"
    output_root = tmp_path / "ws_output"
    data_root.mkdir(parents=True, exist_ok=True)
    templates_root.mkdir(parents=True, exist_ok=True)
    output_root.mkdir(parents=True, exist_ok=True)
    monkeypatch.delenv("RESUME_TAILOR_CALIBRATION_DIR", raising=False)
    for name, value in (
        ("DATA_ROOT", data_root),
        ("TEMPLATES_ROOT", templates_root),
        ("OUTPUT_ROOT", output_root),
        ("CACHE_ROOT", output_root),
        ("DATA_DIR", data_root),
        ("TEMPLATES_DIR", templates_root),
        ("OUTPUT_DIR", output_root),
        ("CACHE_DIR", output_root),
        ("MASTER_RESUME_PATH", data_root / "master_resume.json"),
        ("SETTINGS_PATH", data_root / "settings.json"),
        ("LIBRARIES_PATH", data_root / "libraries.json"),
        ("DEFAULT_TEMPLATE_PATH", templates_root / "main_template.docx"),
        ("BASELINE_TEMPLATE_PATH", templates_root / "original_export.docx"),
        ("TEMPLATE_PROFILE_PATH", templates_root / "template_profile.json"),
        ("TEMPLATE_LIBRARY_DIR", templates_root / "library"),
        ("_ACTIVE_WORKSPACE_ID", None),
    ):
        monkeypatch.setattr(config, name, value)
    return {"data": data_root, "templates": templates_root, "output": output_root}


def _seed_run_for_verify(job_id: str = "verify01") -> Path:
    """Write bullets.json + jd.txt under OUTPUT_DIR/jobs/<id> for verify-claim tests."""
    out_dir = config.OUTPUT_DIR / "jobs" / job_id
    out_dir.mkdir(parents=True, exist_ok=True)
    resume = synthetic_resume()
    bullets = {b.id: b.text for b in resume.all_bullets()}
    (out_dir / "bullets.json").write_text(json.dumps(bullets), encoding="utf-8")
    (out_dir / "jd.txt").write_text(
        "Software Engineer role requiring Python and FastAPI.",
        encoding="utf-8",
    )
    return out_dir
