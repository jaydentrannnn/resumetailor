"""Tests for cover-letter template build and render.

Builds from the same synthetic resume fixture as ``built_template``; no PDF conversion.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from resume_tailor import config
from resume_tailor.document import cover_template, docx_text, render, template_analyze
from resume_tailor.document.template_build import _para_by_id
from resume_tailor.document.template_profile import load_profile
from resume_tailor.pipeline.coverletter import CoverLetter
from tests.fixtures import _docx_bytes, _full_featured_resume, synthetic_resume


@pytest.fixture
def cover_template_paths(tmp_path) -> tuple[Path, Path, Path]:
    """Baseline export, tagged resume template profile source, and cover template out."""
    src = tmp_path / "original_export.docx"
    src.write_bytes(_docx_bytes(_full_featured_resume))
    result = template_analyze.analyze_docx(raw=src.read_bytes())
    assert result.ready, result.issues
    profile_path = tmp_path / "template_profile.json"
    profile_path.write_text(result.suggested_profile.model_dump_json(indent=2), encoding="utf-8")
    cover_dst = tmp_path / "cover_template.docx"
    cover_template.build_cover_template(src, cover_dst, profile=result.suggested_profile)
    return src, cover_dst, profile_path


def test_build_cover_template_produces_placeholders(cover_template_paths):
    """The derived template carries the fixed Jinja placeholder set."""
    _src, cover_dst, _profile = cover_template_paths
    xml = zipfile.ZipFile(cover_dst).read("word/document.xml").decode("utf-8")
    assert "{{ date }}" in xml
    assert "{{ salutation }}" in xml
    assert "{%p for para in paragraphs %}" in xml
    assert "{{ signature }}" in xml


def test_letter_paragraphs_drop_the_bullet_hanging_indent(cover_template_paths):
    """Body donor is a resume bullet, so its indent must be zeroed, not inherited.

    Regression: stripping only ``w:numPr`` left ``w:ind left=720 hanging=360`` on every
    letter paragraph, rendering the whole letter inset half an inch with a negative first
    line instead of flush to the page margin.
    """
    _src, cover_dst, _profile = cover_template_paths
    doc = Document(str(cover_dst))
    tagged = [
        p
        for p, _loc in docx_text.iter_document_paragraphs(doc)
        if p.text.strip() in {"{{ date }}", "{{ line }}", "{{ salutation }}", "{{ para }}"}
    ]
    assert tagged, "no tagged letter paragraphs found"
    for para in tagged:
        fmt = para.paragraph_format
        assert fmt.left_indent == 0, f"{para.text!r} kept a left indent"
        assert fmt.first_line_indent == 0, f"{para.text!r} kept a hanging indent"
        p_pr = para._p.find(qn("w:pPr"))
        assert p_pr is not None
        assert p_pr.find(qn("w:numPr")) is None, f"{para.text!r} kept bullet numbering"


def test_letter_blocks_are_separated_but_address_stays_tight(cover_template_paths):
    """Body blocks get one blank line; the repeating address line and closing do not."""
    _src, cover_dst, _profile = cover_template_paths
    doc = Document(str(cover_dst))
    by_tag = {
        p.text.strip(): p
        for p, _loc in docx_text.iter_document_paragraphs(doc)
        if p.text.strip()
    }
    assert by_tag["{{ para }}"].paragraph_format.space_after > 0
    assert by_tag["{{ date }}"].paragraph_format.space_after > 0
    # The address repeats in a loop, so it cannot space only its last line; the gap
    # before the salutation is carried by the salutation's own space-before instead.
    assert by_tag["{{ line }}"].paragraph_format.space_after == 0
    assert by_tag["{{ salutation }}"].paragraph_format.space_before > 0
    assert by_tag["{{ closing }}"].paragraph_format.space_after == 0
    assert by_tag["{{ signature }}"].paragraph_format.space_after == 0


def test_cover_template_uses_wider_margins(cover_template_paths):
    """Cover letters use ``COVER_MARGIN_INCHES``, not the resume's cram-to-one-page setup."""
    _src, cover_dst, _profile = cover_template_paths
    doc = Document(str(cover_dst))
    margin = doc.sections[0].left_margin
    assert margin == Inches(config.COVER_MARGIN_INCHES)


def test_cover_template_body_uses_relaxed_line_spacing(cover_template_paths):
    """Body paragraphs use 1.15 line spacing instead of the resume bullet's single spacing."""
    _src, cover_dst, _profile = cover_template_paths
    doc = Document(str(cover_dst))
    para = next(
        p
        for p, _loc in docx_text.iter_document_paragraphs(doc)
        if p.text.strip() == "{{ para }}"
    )
    assert para.paragraph_format.line_spacing == 1.15


def test_cover_template_letterhead_carries_section_rule(cover_template_paths):
    """The contact line clones the resume's section-heading bottom border when present."""
    src, cover_dst, profile_path = cover_template_paths
    src_doc = Document(str(src))
    if cover_template._find_rule_donor(src_doc) is None:
        pytest.skip("synthetic baseline has no section-heading rule to clone")
    profile = load_profile(profile_path)
    doc = Document(str(cover_dst))
    contact = _para_by_id(doc, profile.contact.paragraph_id)
    p_pr = contact._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    assert p_bdr is not None
    bottom = p_bdr.find(qn("w:bottom"))
    assert bottom is not None
    assert bottom.get(qn("w:val")) not in (None, "nil", "none")


def test_cover_template_clones_rule_onto_contact_when_baseline_has_one(
    cover_template_paths, tmp_path
):
    """When the baseline carries a section-heading rule, the contact line inherits it."""
    src, _cover_dst, profile_path = cover_template_paths
    profile = load_profile(profile_path)
    ruled_src = tmp_path / "ruled_export.docx"
    doc = Document(str(src))
    donor_para = next(
        p
        for p, _loc in docx_text.iter_document_paragraphs(doc)
        if p.text.strip().upper() == "WORK EXPERIENCES"
    )
    p_pr = donor_para._p.get_or_add_pPr()
    from docx.oxml import OxmlElement

    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    p_pr.insert(0, p_bdr)
    doc.save(str(ruled_src))

    cover_dst = tmp_path / "ruled_cover_template.docx"
    cover_template.build_cover_template(ruled_src, cover_dst, profile=profile)
    out_doc = Document(str(cover_dst))
    contact = _para_by_id(out_doc, profile.contact.paragraph_id)
    contact_p_pr = contact._p.get_or_add_pPr()
    cloned = contact_p_pr.find(qn("w:pBdr"))
    assert cloned is not None
    cloned_bottom = cloned.find(qn("w:bottom"))
    assert cloned_bottom is not None
    assert cloned_bottom.get(qn("w:val")) == "single"


def test_ensure_cover_template_rebuilds_on_builder_version_mismatch(
    cover_template_paths, tmp_path
):
    """A stale ``cover_template.meta.json`` forces a rebuild even when the file exists."""
    src, _cover_dst, profile_path = cover_template_paths
    profile = load_profile(profile_path)
    dst = tmp_path / "cover_template.docx"
    meta = tmp_path / "cover_template.meta.json"
    dst.write_bytes(b"stale")
    meta.write_text('{"builder_version": 0}\n', encoding="utf-8")
    cover_template.ensure_cover_template(src=src, dst=dst, meta_path=meta, profile=profile)
    assert zipfile.is_zipfile(dst)
    assert cover_template._read_builder_version(meta) == cover_template._BUILDER_VERSION


def test_render_cover_letter_against_synthetic_resume(cover_template_paths, tmp_path):
    """A drafted letter renders through the derived template without error."""
    _src, cover_dst, profile_path = cover_template_paths
    resume = synthetic_resume()
    letter = CoverLetter(
        company="Example Corp",
        company_location="Irvine, CA",
        addressee="",
        paragraphs=[
            "I built Python APIs and shipped measurable improvements for users.",
            "The posting emphasizes Python work and I have delivered similar results.",
            "I would welcome a conversation about contributing to your team.",
        ],
        salutation="Dear Hiring Manager,",
        closing="Sincerely,",
        signature=resume.contact.name,
        inside_address=["Example Corp", "Irvine, CA"],
        date="August 29, 2026",
    )
    out = tmp_path / "cover.docx"
    render.render_letter(resume, letter, template=cover_dst, out=out)
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert resume.contact.name in xml
    assert "Dear Hiring Manager," in xml
    assert "Sincerely," in xml
    assert "Example Corp" in xml
