"""JD from a URL or a file (`jd_input`, `/api/jd/*`). No network: fetches are stubbed."""

from __future__ import annotations

import io

import docx
import pytest
from fastapi.testclient import TestClient

from resume_tailor import jd_input
from resume_tailor.apply import fetch_jd
from resume_tailor.web.app import app

POSTING = (
    "Financial Analyst Intern\n\nAbout the role\n"
    + "You will build models in Excel and work with our team on valuation and reporting. " * 8
)


def _pdf(lines: list[str]) -> bytes:
    """A one-page PDF with a real text layer, written by hand (no PDF library needed)."""
    ops = "BT /F1 11 Tf 14 TL 50 750 Td " + " ".join(
        "(" + line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") '"
        for line in lines
    ) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(ops)} >>\nstream\n{ops}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


@pytest.fixture
def fake_fetch(monkeypatch):
    calls: list[dict] = []
    result = {"value": fetch_jd.FetchResult("https://x/final", "greenhouse", POSTING, "api")}

    def _fetch(url, **kwargs):
        calls.append({"url": url, **kwargs})
        return result["value"]

    monkeypatch.setattr(jd_input.fetch_jd, "fetch_jd", _fetch)
    monkeypatch.setattr(jd_input.identity, "resolve_final_url", lambda url: url)
    return calls, result


def test_url_fetch_returns_clean_text(fake_fetch):
    calls, _ = fake_fetch
    got = jd_input.from_url("https://boards.greenhouse.io/acme/jobs/123")
    assert got.text.startswith("Financial Analyst Intern") and got.ats == "greenhouse"
    assert calls[0]["canonical_key"] == "greenhouse:acme:123"
    assert got.warnings == []


@pytest.mark.parametrize(
    "url, code",
    [
        ("ftp://x.com/job", "bad_url"),
        ("not a url", "bad_url"),
        ("https://www.linkedin.com/jobs/view/1", "login_required"),
        ("https://app.joinhandshake.com/jobs/1", "login_required"),
    ],
)
def test_url_rejections_need_no_network(fake_fetch, url, code):
    calls, _ = fake_fetch
    with pytest.raises(jd_input.JdInputError) as err:
        jd_input.from_url(url)
    assert err.value.code == code and calls == []


def test_closed_and_unreadable_pages(fake_fetch):
    _, result = fake_fetch
    result["value"] = fetch_jd.FetchResult("u", "other", "", "failed", error="Client error '404 Not Found'")
    with pytest.raises(jd_input.JdInputError, match="no longer exists"):
        jd_input.from_url("https://jobs.example.com/1")
    result["value"] = fetch_jd.FetchResult("u", "other", "Loading…", "failed", error="too short")
    with pytest.raises(jd_input.JdInputError) as err:
        jd_input.from_url("https://jobs.example.com/1")
    assert err.value.code == "unreadable"


def test_warnings_for_closed_long_and_non_english(fake_fetch):
    _, result = fake_fetch
    closed = "This job is no longer accepting applications.\n" + POSTING
    result["value"] = fetch_jd.FetchResult("u", "other", closed, "http")
    assert "closed" in jd_input.from_url("https://jobs.example.com/1").warnings[0]

    huge = POSTING * 200
    result["value"] = fetch_jd.FetchResult("u", "other", huge, "http")
    got = jd_input.from_url("https://jobs.example.com/1")
    assert len(got.text) == jd_input.MAX_CHARS and "cut off" in got.warnings[0]

    german = "Wir suchen eine engagierte Person für unser Team in Berlin. " * 20
    result["value"] = fetch_jd.FetchResult("u", "other", german, "http")
    assert "English" in jd_input.from_url("https://jobs.example.com/1").warnings[0]


def test_short_text_is_an_error():
    with pytest.raises(jd_input.JdInputError) as err:
        jd_input.from_file("jd.txt", b"Apply now")
    assert err.value.code == "too_short"


def test_text_html_docx_and_pdf_files():
    assert jd_input.from_file("jd.txt", POSTING.encode("utf-8-sig")).text.startswith("Financial")
    assert jd_input.from_file("jd.txt", POSTING.encode("cp1252")).source == "txt"
    html = f"<html><body><nav>Menu</nav><main><p>{POSTING}</p></main></body></html>"
    assert "valuation" in jd_input.from_file("jd.html", html.encode()).text

    document = docx.Document()
    for line in POSTING.splitlines():
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Location", "New York"
    buf = io.BytesIO()
    document.save(buf)
    got = jd_input.from_file("jd.docx", buf.getvalue())
    assert "Location | New York" in got.text

    sentence = "You will build models in Excel and work on valuation and reporting."
    pdf = _pdf(["Financial Analyst Intern"] + [sentence] * 6)
    assert "valuation" in jd_input.from_file("posting.PDF", pdf).text


@pytest.mark.parametrize(
    "name, raw, code",
    [
        ("jd.exe", b"x", "file_type"),
        ("jd.txt", b"x" * (jd_input.MAX_FILE_BYTES + 1), "too_large"),
        ("jd.pdf", b"not a pdf", "unreadable"),
        ("jd.docx", b"not a zip", "unreadable"),
    ],
)
def test_bad_files(name, raw, code):
    with pytest.raises(jd_input.JdInputError) as err:
        jd_input.from_file(name, raw)
    assert err.value.code == code


def test_image_only_pdf_is_reported_as_scanned():
    with pytest.raises(jd_input.JdInputError) as err:
        jd_input.from_file("scan.pdf", _pdf([]))
    assert err.value.code == "scanned"


def test_routes(fake_fetch):
    with TestClient(app) as c:
        ok = c.post("/api/jd/fetch", json={"url": "https://boards.greenhouse.io/acme/jobs/1"})
        assert ok.status_code == 200 and ok.json()["ats"] == "greenhouse"
        bad = c.post("/api/jd/fetch", json={"url": "https://linkedin.com/jobs/1"})
        assert bad.status_code == 422 and bad.json()["error"] == "login_required"
        up = c.post("/api/jd/extract-file", files={"file": ("jd.txt", POSTING.encode(), "text/plain")})
        assert up.status_code == 200 and up.json()["source"] == "txt"
        wrong = c.post("/api/jd/extract-file", files={"file": ("jd.exe", b"x", "application/x")})
        assert wrong.json()["error"] == "file_type"
